"""Real offline exports, provenance and preview-mode regressions."""
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader

from rpa_docs import pipeline
from rpa_docs.evidence_documents import figure_catalog, source_timeline, attach_source, export_interchange, wants_previews
from rpa_docs.documentation_profiles import PROFILES
from rpa_docs.schemas import JobCreate
from rpa_docs.orchestrator import rebuild


def fixture(root, profile="processo", mode="documentacao"):
    (root / "evidence").mkdir(exist_ok=True)
    steps = []
    for index, color in [(1, "#f0eaff"), (2, "#e7f4ef")]:
        image = Image.new("RGB", (960, 540), color)
        draw = ImageDraw.Draw(image)
        draw.text((48, 40), f"QA tela {index} - Pedido 0042", fill="#130d29", font_size=32)
        draw.rectangle((48, 100, 900, 460), outline="#571ee6", width=3)
        draw.text((80, 150), "Status: " + ("Em revisao" if index == 1 else "Confirmado"), fill="#130d29", font_size=28)
        name = f"tela-{index}.png"
        image.save(root / "evidence" / name)
        steps.append({"frame_index": index, "time": 6 if index == 1 else 22,
                      "timecode": "00:00:06" if index == 1 else "00:00:22",
                      "image": name, "title": "Conferir pedido" if index == 1 else "Confirmar pedido",
                      "action": "Demonstração sintética de QA.", "evidence": f"Tela {index} mostra o status do pedido.",
                      "source_name": "sessao-a.mp4" if index == 1 else "sessao-b.mp4",
                      "source_timecode": "00:00:06" if index == 1 else "00:00:02"})
    (root / "analise-visual.json").write_text(json.dumps(steps), encoding="utf-8")
    (root / "transcricao.json").write_text("[]", encoding="utf-8")
    context = {"perfil": profile, "modoEntrega": mode, "processo": "Conferência de pedidos",
               "publico": ["equipe"], "entregaveis": ["fluxo_atual"] if profile == "processo" else PROFILES[profile]["suggestions"]}
    if profile == "processo":
        context.update(cliente="Ambiente sintético", solucoes=["rpa"], etapaDestino="spec_rpa")
    job = JobCreate(contexto=context, source_path=str(root / "original.mp4"), api_approved=True).model_dump()
    job["id"] = "qa-static"
    document = pipeline.fallback_documentation(job, steps, "Dados sintéticos para QA.")
    if profile != "processo":
        for section in document["sections"]:
            section["items"] = [{"title": "Conferir e confirmar", "text": "Conteúdo sintético relacionado às duas telas.",
                                 "evidence_refs": [1, 2], "timecode": "00:00:06"}]
    return job, document, steps


class StaticEvidenceTests(unittest.TestCase):
    def test_source_local_time_crosses_recording_boundaries(self):
        timeline = source_timeline([Path("parte1.mp4"), Path("parte2.mp4")], [20, 30])
        item = attach_source({"time": 22}, timeline)
        self.assertEqual(item["source_name"], "parte2.mp4")
        self.assertEqual(item["source_timecode"], "00:00:02")
        self.assertEqual(attach_source({"time": 20}, timeline)["source_id"], "source-002")

    def test_path_escape_and_corrupt_images_are_never_embedded(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "evidence").mkdir()
            Image.new("RGB", (20, 20)).save(root / "outside.png")
            (root / "evidence" / "corrupt.png").write_bytes(b"invalid")
            catalog = figure_catalog(root, [{"frame_index": 1, "image": "../outside.png"},
                                           {"frame_index": 2, "image": "corrupt.png"}])
            self.assertTrue(all(figure.path is None for figure in catalog.values()))

    def test_six_profiles_embed_the_correct_images_in_pdf_and_word(self):
        for profile in PROFILES:
            with self.subTest(profile=profile), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                job, document, steps = fixture(root, profile)
                pipeline.write_report(job, steps, root, Path("source.mp4"), document, [])
                stem = "documentacao-processo" if profile == "processo" else "documentacao"
                word = Document(root / (stem + ".docx"))
                self.assertEqual(len(word.inline_shapes), 2)
                with zipfile.ZipFile(root / (stem + ".docx")) as archive:
                    self.assertEqual(len([name for name in archive.namelist() if name.startswith("word/media/")]), 2)
                    self.assertNotIn('TargetMode="External"', archive.read("word/_rels/document.xml.rels").decode())
                pdf = PdfReader(root / (stem + ".pdf"))
                self.assertEqual(sum(len(page.images) for page in pdf.pages), 2)
                text = "\n".join(page.extract_text() for page in pdf.pages)
                self.assertIn("sessao-b.mp4", text)
                self.assertIn("00:00:02", text)
                html = (root / "relatorio.html").read_text(encoding="utf-8")
                self.assertIn("data:image/png;base64,", html)
                self.assertNotIn("<video", html)
                pipeline.create_package(root)
                with zipfile.ZipFile(root / "entrega-completa.zip") as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertIn("manifesto-documentacao.json", archive.namelist())
                    self.assertIn("evidence/tela-1.png", archive.namelist())
                    self.assertIn("evidence/tela-2.png", archive.namelist())
                    self.assertFalse(any(name.endswith(".mp4") for name in archive.namelist()))

    def test_interchange_contains_assets_hashes_and_no_machine_path(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job, document, steps = fixture(root)
            manifest = export_interchange(job, document, steps, root)
            self.assertEqual(len(manifest["evidence"]), 2)
            self.assertTrue(all(len(item["sha256"]) == 64 for item in manifest["evidence"]))
            self.assertNotIn(str(root), json.dumps(manifest))
            with zipfile.ZipFile(root / "intercambio-documentacao.zip") as archive:
                self.assertEqual(set(archive.namelist()), {"manifesto-documentacao.json", "evidence/tela-1.png", "evidence/tela-2.png"})
                self.assertIsNone(archive.testzip())

    def test_missing_print_has_an_explicit_warning_not_a_substitute(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job, document, steps = fixture(root)
            (root / "evidence" / "tela-2.png").unlink()
            pipeline.write_report(job, steps, root, Path("source.mp4"), document, [])
            word = Document(root / "documentacao-processo.docx")
            self.assertEqual(len(word.inline_shapes), 1)
            self.assertIn("Evidência 2 indisponível", "\n".join(p.text for p in word.paragraphs))

    def test_large_step_text_and_long_caption_paginate(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job, document, steps = fixture(root)
            steps[0]["title"] = "Título longo " * 80
            steps[0]["evidence"] = "Descrição confirmada " * 200
            document["process_flow"][0]["description"] = "Explicação da operação. " * 600
            (root / "analise-visual.json").write_text(json.dumps(steps), encoding="utf-8")
            pipeline.write_report(job, steps, root, Path("source.mp4"), document, [])
            pdf = PdfReader(root / "documentacao-processo.pdf")
            self.assertGreater(len(pdf.pages), 2)
            self.assertEqual(sum(len(page.images) for page in pdf.pages), 2)

    def test_document_mode_never_calls_video_renderer(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job, document, steps = fixture(root)
            def transcribe(video, target):
                return [{"start": 6, "end": 8, "text": "Conferir pedido."}]
            with patch.object(pipeline, "prepare_source", return_value=(Path("source.mp4"), [Path("source.mp4")])), \
                 patch.object(pipeline, "probe_duration", return_value=30), \
                 patch.object(pipeline, "extract_evidence", return_value=steps), \
                 patch.object(pipeline, "transcribe", side_effect=transcribe), \
                 patch.object(pipeline, "analyze", return_value=steps), \
                 patch.object(pipeline, "synthesize_documentation", return_value=document), \
                 patch.object(pipeline, "create_job_previews") as renderer:
                result = pipeline.process(job, root, lambda *args: None)
            renderer.assert_not_called()
            self.assertFalse(result["preview"]["requested"])
            self.assertEqual(result["warnings"], [])

    def test_legacy_requests_still_request_previews(self):
        self.assertTrue(wants_previews({}))
        self.assertTrue(wants_previews({"contexto": {"perfil": "processo"}}))
        self.assertFalse(wants_previews({"contexto": {"modoEntrega": "documentacao"}}))

    def test_malformed_timeline_degrades_without_reassigning_a_source(self):
        original = {"time": 22, "source_name": "confirmed.mp4"}
        self.assertEqual(attach_source(original, [None, {}, {"offset": "bad", "duration": 30}]), original)

    def test_operational_chapter_heading_stays_with_first_illustrated_step(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job, document, steps = fixture(root)
            # Fill the preceding page enough to trigger the chapter's page break.
            document["scope"]["in_scope"] = ["Conferência de um pedido no sistema." for _ in range(4)]
            pipeline.write_report(job, steps, root, Path("source.mp4"), document, [])
            pages = PdfReader(root / "documentacao-processo.pdf").pages
            chapter = next(page.extract_text() for page in pages if "7. Fluxo operacional" in page.extract_text())
            self.assertIn("1. Conferir pedido", chapter)

    def test_document_only_rebuild_does_not_require_missing_original_video(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            job, document, steps = fixture(root)
            pipeline.write_report(job, steps, root, Path("missing.mp4"), document, [])
            job["result_json"] = {"video": str(root / "missing.mp4"), "source_files": [str(root / "missing.mp4")]}
            with patch("rpa_docs.orchestrator.synthesize_documentation", return_value=document), patch("rpa_docs.orchestrator.create_job_previews") as renderer:
                result = rebuild(job, root, lambda *args: None)
            renderer.assert_not_called()
            self.assertNotIn("previews separados", result["deliverables"])
            self.assertEqual(len(Document(root / "documentacao-processo.docx").inline_shapes), 2)

    def test_real_ffmpeg_audio_decode_uses_waveform_not_incompatible_pyav_adapter(self):
        import numpy as np
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "audio.wav"
            subprocess.run([pipeline.FFMPEG, "-nostdin", "-v", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=1", "-ac", "2", str(source)], check=True)
            model = SimpleNamespace(transcribe=lambda waveform, **kwargs: (iter([SimpleNamespace(start=0, end=1, text=" Teste ")]), {}))
            from unittest.mock import Mock
            model.transcribe = Mock(wraps=model.transcribe)
            with patch.dict(sys.modules, {"faster_whisper": SimpleNamespace(WhisperModel=lambda *args, **kwargs: model)}):
                result = pipeline.transcribe(source, root)
            waveform = model.transcribe.call_args.args[0]
            self.assertIsInstance(waveform, np.ndarray)
            self.assertEqual(waveform.dtype, np.float32)
            self.assertEqual(waveform.size, 16000)
            self.assertEqual(result[0]["text"], "Teste")


if __name__ == "__main__":
    unittest.main()
