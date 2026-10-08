"""Offline checks for portable, provenance-aware structured deliverables."""

import importlib.util
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document

from rpa_docs.pipeline import fallback_documentation, report_styles, write_report


class DeliverablesQATests(unittest.TestCase):
    def setUp(self):
        self.job = {
            "title": "[QA sintético] Gestão de pedidos — Trevo",
            "audience": "Equipe de RPA",
            "detail_level": "operacional",
            "process_context": "Documentar pedidos da Trevo. Remover conversas sem relação; manter exceções.",
        }
        self.steps = [{
            "frame_index": 1, "time": 6, "timecode": "00:00:06",
            "title": "Registrar o pedido", "action": "Demonstração sintética; não representa reunião real.",
            "system": "Ambiente de teste", "image": "evidencia-0001.jpg",
        }]
        self.warning = "Dados sintéticos de QA, sem inferência de IA."

    def test_fallback_does_not_promote_cutting_instructions_to_business_facts(self):
        document = fallback_documentation(self.job, self.steps, self.warning)
        self.assertNotIn(self.job["process_context"], document["objective"])
        self.assertNotIn(self.job["process_context"], document["executive_summary"])
        self.assertIn("não confirmado", document["objective"])
        self.assertIn(self.warning, " ".join(document["limitations"]))
        self.assertEqual(document["process_flow"][0]["evidence_refs"], [1])

    def test_exported_guidance_and_warnings_accompany_each_standalone_document(self):
        document = fallback_documentation(self.job, self.steps, self.warning)
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary)
            write_report(self.job, self.steps, workspace, Path("synthetic.mp4"), document)
            for name in ("procedimento-operacional.md", "requisitos-rpa.md", "relatorio.html"):
                with self.subTest(name=name):
                    text = (workspace / name).read_text(encoding="utf-8")
                    self.assertIn("Guia recebido para esta análise", text)
                    self.assertIn(self.job["process_context"], text)
                    self.assertIn(self.warning, text)
                    self.assertIn("não equivale a um objetivo de negócio confirmado", text)
            word = Document(workspace / "documentacao-processo.docx")
            paragraphs = "\n".join(p.text for p in word.paragraphs)
            self.assertIn(self.warning, paragraphs)
            self.assertIn("Guia recebido para esta análise", paragraphs)
            self.assertEqual(word.paragraphs[1].style.name, "Title")
            self.assertEqual(word.styles["Normal"].font.name, "Arial")
            self.assertEqual(str(word.styles["Normal"].font.color.rgb), "060315")
            self.assertEqual(len([p for p in word.paragraphs if p.style.name == "Heading 1"]), 13)
            with zipfile.ZipFile(workspace / "documentacao-processo.docx") as package:
                self.assertIn('w:instr="PAGE"', package.read("word/footer1.xml").decode("utf-8"))

    def test_html_styles_bundle_offline_branding_fonts_and_focus(self):
        css = report_styles()
        self.assertEqual(css.count("data:font/woff;base64,"), 3)
        self.assertIn("#571ee6", css)
        self.assertIn("#060315", css)
        self.assertIn(":focus-visible", css)
        self.assertNotIn("https://", css)

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is an optional QA reader")
    def test_pdf_warnings_remain_with_heading_and_long_titles_stay_in_footer(self):
        from pypdf import PdfReader

        document = fallback_documentation(self.job, self.steps, self.warning)
        self.job["title"] = "Título extenso de conferência de documentação " * 4
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary)
            write_report(self.job, self.steps, workspace, Path("synthetic.mp4"), document)
            reader = PdfReader(workspace / "documentacao-processo.pdf")
            pages = [page.extract_text() for page in reader.pages]
            warning_pages = [page for page in pages if self.warning in page]
            self.assertEqual(len(warning_pages), 1)
            self.assertIn("13. Limitações da análise", warning_pages[0])
            self.assertTrue(all("página " in page for page in pages))
            self.assertIn("Guia recebido para esta análise", "\n".join(pages))
            self.assertIn("Compatibilidade: Word em Arial e PDF em Helvetica", "\n".join(pages))
            footer_positions = []
            for page in reader.pages:
                page.extract_text(visitor_text=lambda text, cm, tm, font, size: footer_positions.append(tm[4]) if "página " in text and size == 8.5 else None)
            self.assertEqual(len(footer_positions), len(pages))
            self.assertTrue(all(x >= 71.9 for x in footer_positions))


if __name__ == "__main__":
    unittest.main()
