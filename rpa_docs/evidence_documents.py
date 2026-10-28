"""Shared, local evidence figures for document exporters and As Is interchange."""
from __future__ import annotations

import hashlib
import json
import math
import zipfile
from dataclasses import dataclass
from html import escape
from pathlib import Path

DOCUMENT_ONLY = "documentacao"
WITH_PREVIEWS = "documentacao_previews"
INTERCHANGE_FILES = {"fontes.json", "manifesto-documentacao.json", "intercambio-documentacao.zip"}


def wants_previews(job):
    return (job.get("contexto") or {}).get("modoEntrega", WITH_PREVIEWS) != DOCUMENT_ONLY


def clock(seconds):
    try:
        value = max(0, int(float(seconds)))
    except (ValueError, TypeError, OverflowError):
        return ""
    return f"{value // 3600:02}:{value // 60 % 60:02}:{value % 60:02}"


def seconds(value):
    try:
        if isinstance(value, str) and ":" in value:
            parts = [float(part) for part in value.split(":")]
            return sum(part * 60 ** index for index, part in enumerate(reversed(parts)))
        return float(value)
    except (ValueError, TypeError):
        return None


def source_timeline(videos, durations):
    result, offset = [], 0.0
    for index, (video, duration) in enumerate(zip(videos, durations), 1):
        result.append({"source_id": f"source-{index:03}", "name": Path(video).name,
                       "duration": round(duration, 3), "offset": round(offset, 3)})
        offset += duration
    return result


def attach_source(item, timeline):
    moment = seconds(item.get("time", item.get("timecode")))
    if moment is None or not math.isfinite(moment):
        return dict(item)
    for source in timeline:
        if not isinstance(source, dict):
            continue
        offset, duration = seconds(source.get("offset")), seconds(source.get("duration"))
        if offset is None or duration is None or not math.isfinite(offset + duration) or duration <= 0:
            continue
        if offset <= moment < offset + duration:
            return {**item, "source_id": str(source.get("source_id") or ""), "source_name": Path(str(source.get("name") or "")).name,
                    "source_time": round(moment - offset, 3),
                    "source_timecode": clock(moment - offset)}
    return dict(item)


@dataclass
class Figure:
    ref: str
    path: Path | None = None
    title: str = ""
    description: str = ""
    source: str = ""
    timecode: str = ""
    uncertainty: str = ""
    reason: str = ""
    width: int = 0
    height: int = 0

    @property
    def caption(self):
        return f"Evidência {self.ref} - {self.title or 'Tela demonstrada'}"

    @property
    def provenance(self):
        return " | ".join(value for value in (self.source, self.timecode) if value) or "Origem não identificada"

    @property
    def unavailable(self):
        return f"Evidência {self.ref} indisponível: {self.reason or 'imagem não encontrada'}. Revise a fonte."

    def size(self, max_width, max_height):
        factor = min(max_width / self.width, max_height / self.height)
        return self.width * factor, self.height * factor


def read_json(path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return default


def figure_catalog(workspace, steps=None):
    from PIL import Image
    workspace = Path(workspace).resolve()
    report = read_json(workspace / "relatorio.json", {})
    if not isinstance(report, dict):
        report = {}
    if steps is None:
        steps = read_json(workspace / "analise-visual.json", report.get("evidence_steps", []))
    if not isinstance(steps, list):
        steps = []
    timeline = read_json(workspace / "fontes.json", [])
    if not isinstance(timeline, list):
        timeline = []
    result = {}
    for step in steps:
        if not isinstance(step, dict) or step.get("frame_index") is None:
            continue
        step = attach_source(step, timeline) if timeline else step
        ref = str(step["frame_index"])
        figure = Figure(ref=ref, title=str(step.get("title") or ""),
                        description=str(step.get("evidence") or step.get("action") or ""),
                        source=str(step.get("source_name") or Path(str(report.get("source") or "")).name),
                        timecode=str(step.get("source_timecode") or step.get("timecode") or clock(step.get("time"))),
                        uncertainty=str(step.get("uncertainty") or ""))
        try:
            relative = Path(str(step.get("image") or ""))
            root = workspace / "evidence"
            candidate = (root / relative).resolve()
            if not str(relative) or str(relative) == "." or relative.is_absolute() or not candidate.is_relative_to(root.resolve()):
                raise ValueError("caminho inválido")
            if candidate.suffix.lower() not in {".jpg", ".jpeg", ".png"}:
                raise ValueError("formato de imagem não suportado")
            with Image.open(candidate) as img:
                if img.format not in {"JPEG", "PNG"} or max(img.size) > 24000:
                    raise ValueError("imagem fora dos limites")
                figure.width, figure.height = img.size
                img.verify()
            figure.path = candidate
        except (OSError, ValueError, Image.DecompressionBombError):
            figure.reason = "imagem ausente, inválida ou fora da pasta de evidências"
        result[ref] = figure
    return result


def figures(refs, catalog):
    return [catalog.get(ref, Figure(ref, reason="referência sem imagem associada"))
            for ref in dict.fromkeys(str(value) for value in (refs or []))]


def word_figures(word, refs, catalog, seen):
    from docx.shared import Inches, Pt
    items = figures(refs, catalog)
    if not items:
        word.add_paragraph("Sem evidência visual associada. Valide a descrição com as falas da gravação.")
    section = word.sections[-1]
    width = (section.page_width - section.left_margin - section.right_margin) / 914400
    for figure in items:
        if figure.path is None:
            word.add_paragraph(figure.unavailable)
        elif figure.ref in seen:
            word.add_paragraph(f"Ver evidência {figure.ref}, apresentada anteriormente. {figure.provenance}.")
        else:
            w, h = figure.size(width, 3.65)
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(5)
            shape = paragraph.add_run().add_picture(str(figure.path), width=Inches(w), height=Inches(h))
            shape._inline.docPr.set("descr", figure.description or figure.title)
            caption = word.add_paragraph(figure.caption)
            caption.paragraph_format.space_after = Pt(4)
            for run in caption.runs:
                run.bold = True
                run.font.size = Pt(9)
            word.add_paragraph(figure.provenance)
            if figure.description:
                word.add_paragraph(figure.description)
            if figure.uncertainty:
                word.add_paragraph("A validar: " + figure.uncertainty)
            seen.add(figure.ref)


def pdf_figures(refs, catalog, seen, max_width=480):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Image, KeepTogether, Paragraph, Spacer
    note = ParagraphStyle("figure-note", fontName="Helvetica", fontSize=9, leading=13, spaceAfter=6)
    caption = ParagraphStyle("figure-caption", parent=note, fontName="Helvetica-Bold")
    result = []
    items = figures(refs, catalog)
    if not items:
        result.append(Paragraph("Sem evidência visual associada. Valide a descrição com as falas da gravação.", note))
    for figure in items:
        if figure.path is None:
            result.append(Paragraph(escape(figure.unavailable), note))
        elif figure.ref in seen:
            result.append(Paragraph(escape(f"Ver evidência {figure.ref}, apresentada anteriormente. {figure.provenance}."), note))
        else:
            width, height = figure.size(max_width, 230)
            # The short caption stays with the image; long evidence descriptions
            # flow naturally instead of forming an unbreakable full-page block.
            short = figure.caption[:180]
            result += [Spacer(1, 8), KeepTogether([
                Image(str(figure.path), width=width, height=height, hAlign="LEFT"),
                Spacer(1, 5), Paragraph(escape(short), caption),
            ]), Paragraph(escape(figure.provenance), note)]
            if figure.description:
                result.append(Paragraph(escape(figure.description), note))
            if figure.uncertainty:
                result.append(Paragraph(escape("A validar: " + figure.uncertainty), note))
            seen.add(figure.ref)
    return result


def html_figures(refs, catalog):
    import base64
    result = []
    for figure in figures(refs, catalog):
        if figure.path is None:
            result.append(f'<p class="empty">{escape(figure.unavailable)}</p>')
            continue
        mime = "image/png" if figure.path.suffix.lower() == ".png" else "image/jpeg"
        encoded = base64.b64encode(figure.path.read_bytes()).decode("ascii")
        result.append(f'<figure class="document-figure"><img src="data:{mime};base64,{encoded}" alt="{escape(figure.description or figure.title)}">'
                      f'<figcaption><strong>{escape(figure.caption)}</strong><br>{escape(figure.provenance)}</figcaption>'
                      f'<p>{escape(figure.description)}</p></figure>')
    return "".join(result) or '<p class="empty">Sem evidência visual associada a este item.</p>'


def export_interchange(job, document, steps, workspace):
    catalog = figure_catalog(workspace, steps)
    references = []
    for section in document.get("sections", []):
        for item in section.get("items", []):
            references.extend(item.get("evidence_refs") or [])
    for item in document.get("process_flow", []):
        references.extend(item.get("evidence_refs") or [])
    for item in document.get("business_rules", []):
        references.extend(item.get("evidence_refs") or [])
    evidence = []
    for figure in figures(references, catalog):
        entry = {"id": figure.ref, "title": figure.title, "description": figure.description,
                 "source": figure.source, "timecode": figure.timecode, "uncertainty": figure.uncertainty,
                 "available": figure.path is not None}
        if figure.path:
            entry.update(path=figure.path.relative_to(workspace.resolve()).as_posix(),
                         sha256=hashlib.sha256(figure.path.read_bytes()).hexdigest(),
                         width=figure.width, height=figure.height)
        else:
            entry["warning"] = figure.unavailable
        evidence.append(entry)
    manifest = {"schema_version": 1, "document_id": job.get("id") or "document",
                "title": job.get("title"), "profile": (job.get("contexto") or {}).get("perfil", "processo"),
                "documentation": document, "sources": read_json(workspace / "fontes.json", []),
                "evidence": evidence, "review_status": "pending"}
    (workspace / "manifesto-documentacao.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary = workspace / "intercambio-documentacao.zip.tmp"
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(workspace / "manifesto-documentacao.json", "manifesto-documentacao.json")
        for entry in evidence:
            if entry.get("path"):
                archive.write(workspace / entry["path"], entry["path"])
    temporary.replace(workspace / "intercambio-documentacao.zip")
    return manifest
