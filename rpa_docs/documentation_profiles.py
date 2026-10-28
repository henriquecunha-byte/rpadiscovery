"""Profile contracts and exports. Historical jobs keep the process schema."""
import json
import re
from html import escape

from .context_model import RULES, LABELS
from .evidence_documents import figure_catalog, figures, word_figures, pdf_figures, html_figures

PROFILES = RULES["profiles"]
GENERAL_FILES = {value["filename"] for key, value in PROFILES.items() if key != "processo"}
GENERAL_EXPORTS = GENERAL_FILES | {"documentacao.docx", "documentacao.pdf", "documentacao.json"}


def profile_id(job):
    context = job.get("contexto") or {}
    if hasattr(context, "model_dump"):
        context = context.model_dump()
    return context.get("perfil", "processo")


def general_job(job):
    return profile_id(job) in PROFILES and profile_id(job) != "processo"


def requested_sections(job):
    context = job.get("contexto") or {}
    profile = PROFILES[profile_id(job)]
    selected = context.get("entregaveis", profile.get("suggestions", []))
    return [key for key in LABELS["entregaveis"] if key in selected and key in profile.get("available", [])]


def general_contract(job):
    profile = PROFILES[profile_id(job)]
    shape = {
        "schema_version": 2, "profile": profile_id(job),
        "executive_summary": "Síntese adequada ao objetivo e público",
        "objective": "Objetivo confirmado pelos insumos",
        "scope": {"in_scope": [], "out_of_scope": []},
        "sections": [{"id": key, "title": LABELS["entregaveis"][key],
                      "items": [{"title": "Título específico", "text": "Explicação estruturada, com contexto e evidência",
                                 "evidence_refs": [], "timecode": "00:00:00"}]} for key in requested_sections(job)],
        "open_questions": [], "limitations": [],
        "preview_moments": [{"start": "00:00:00", "end": "00:00:10", "title": "Assunto", "reason": "Relação com o pedido"}],
    }
    return (f'PERFIL DE DOCUMENTAÇÃO: {profile["label"]}\n{profile["focus"]}\n'
            'Use exclusivamente as seções solicitadas abaixo, com itens completos e explicações; não basta listar frases da transcrição. '
            'Resumo, objetivo, escopo, pontos a validar e limitações são comuns a todos os perfis. '
            'Não inclua campos ou seções de processo/RPA por padrão. Para um item sem evidência visual, use evidence_refs vazio. '
            'timecode deve existir na transcrição ou ficar vazio. Se uma seção não tiver evidência, mantenha items vazio e explique a lacuna nas limitações. '
            'Não transforme hipóteses em fatos, sugestões em decisões ou a ausência de evidência em uma conclusão. '
            'No perfil Personalizado, adapte os títulos e a organização dos itens às instruções adicionais.\n'
            'Retorne JSON puro exatamente com esta estrutura:\n' + json.dumps(shape, ensure_ascii=False))


def fallback_general(job, warning=""):
    return {
        "schema_version": 2, "profile": profile_id(job),
        "executive_summary": "Registro preliminar. A síntese do conteúdo não está disponível; consulte as evidências e valide antes de compartilhar.",
        "objective": "Objetivo não confirmado nos insumos disponíveis.",
        "scope": {"in_scope": [], "out_of_scope": []},
        "sections": [{"id": key, "title": LABELS["entregaveis"][key], "items": []} for key in requested_sections(job)],
        "open_questions": ["Validar o conteúdo solicitado com a pessoa responsável pela gravação."],
        "limitations": ["A síntese estruturada não pôde ser confirmada."] + ([warning] if warning else []),
        "preview_moments": [],
    }


def normalize_general(document, job, steps):
    result = fallback_general(job)
    if not isinstance(document, dict):
        return result
    def string(value):
        return value.strip() if isinstance(value, str) else ""
    def strings(value):
        return [item.strip() for item in value if isinstance(item, str) and item.strip()] if isinstance(value, list) else []
    for key in ("executive_summary", "objective"):
        result[key] = string(document.get(key)) or result[key]
    scope = document.get("scope") if isinstance(document.get("scope"), dict) else {}
    result["scope"] = {key: strings(scope.get(key)) for key in ("in_scope", "out_of_scope")}
    refs = {str(step["frame_index"]): step["frame_index"] for step in steps if step.get("frame_index") is not None}
    sections = document.get("sections") if isinstance(document.get("sections"), list) else []
    by_id = {section.get("id"): section for section in sections if isinstance(section, dict) and isinstance(section.get("id"), str)}
    for section in result["sections"]:
        raw_items = by_id.get(section["id"], {}).get("items", [])
        for item in raw_items if isinstance(raw_items, list) else []:
            if not isinstance(item, dict) or not string(item.get("text")):
                continue
            raw_refs = item.get("evidence_refs") if isinstance(item.get("evidence_refs"), list) else []
            timecode = string(item.get("timecode"))
            section["items"].append({
                "title": string(item.get("title")), "text": string(item.get("text")),
                "evidence_refs": list(dict.fromkeys(refs[str(ref)] for ref in raw_refs if str(ref) in refs)),
                "timecode": timecode if re.fullmatch(r"\d{2,3}:[0-5]\d:[0-5]\d", timecode) else "",
            })
    result["open_questions"] = strings(document.get("open_questions"))
    result["limitations"] = strings(document.get("limitations"))
    for section in result["sections"]:
        if not section["items"]:
            result["limitations"].append(f'Sem conteúdo confirmado para: {section["title"]}.')
    moments = document.get("preview_moments")
    result["preview_moments"] = [dict(item) for item in moments if isinstance(item, dict)] if isinstance(moments, list) else []
    return result


def item_count(document):
    if document.get("schema_version") == 2:
        return sum(len(section.get("items", [])) for section in document.get("sections", []))
    return len(document.get("process_flow", []))


def document_catalog(job):
    profile = PROFILES.get(profile_id(job), PROFILES["processo"])
    stem = "documentacao" if general_job(job) else "documentacao-processo"
    base = [(stem + ".docx", profile["document"], "Word"),
            (stem + ".pdf", profile["document"], "PDF")]
    base += [(profile["filename"], profile["document"], "Markdown")] if general_job(job) else [
        ("procedimento-operacional.md", "Procedimento operacional", "Markdown"),
        ("requisitos-rpa.md", "Requisitos de automação", "Markdown")]
    return base + [("intercambio-documentacao.zip", "Dados e prints para integração", "ZIP"),
                   ("matriz-evidencias.csv", "Matriz de evidências", "CSV"),
                   (stem + ".json", "Documentação estruturada", "JSON"),
                   ("roteiro-cortes.json", "Roteiro dos cortes", "JSON"),
                   ("transcricao.json", "Transcrição de referência", "JSON")]


def document_blocks(job, document, with_figures=False):
    """One semantic source for Word, PDF, Markdown and HTML."""
    blocks = [("h1", "Resumo"), ("p", document["executive_summary"]),
              ("h1", "Objetivo"), ("p", document["objective"])]
    for key, title in (("in_scope", "Escopo incluído"), ("out_of_scope", "Fora do escopo")):
        if document["scope"].get(key):
            blocks.append(("h1", title))
            blocks.extend(("bullet", value) for value in document["scope"][key])
    for section in document["sections"]:
        blocks.append(("h1", section["title"]))
        if not section["items"]:
            blocks.append(("p", "Nenhum conteúdo confirmado para este item. Consulte as limitações e valide com a pessoa responsável."))
        for item in section["items"]:
            if item["title"]:
                blocks.append(("h2", item["title"]))
            blocks.append(("p", item["text"]))
            references = []
            if item["timecode"]:
                references.append("Momento da gravação: " + item["timecode"])
            if item["evidence_refs"]:
                references.append("Evidências: " + ", ".join(map(str, item["evidence_refs"])))
            if references:
                blocks.append(("note", " · ".join(references)))
            if with_figures:
                blocks.append(("figure", item.get("evidence_refs", [])))
    for key, title in (("open_questions", "Pontos a validar"), ("limitations", "Limitações da análise")):
        if document.get(key):
            blocks.append(("h1", title))
            blocks.extend(("bullet", value) for value in document[key])
    blocks += [("h1", "Guia recebido para esta análise"), ("p", job.get("process_context") or "Não informado."),
               ("note", "O guia é uma orientação do pedido, não uma comprovação dos fatos. Revise as evidências antes de compartilhar.")]
    return blocks


def write_general_markdown(job, document, workspace):
    profile = PROFILES[profile_id(job)]
    prefixes = {"h1": "## ", "h2": "### ", "bullet": "- ", "p": "", "note": "> "}
    lines = [f'# {profile["document"]} — {job.get("title", "Análise")}', f'Público: {job.get("audience") or "Equipe"}']
    catalog = figure_catalog(workspace)
    for kind, value in document_blocks(job, document, with_figures=True):
        if kind == "figure":
            for figure in figures(value, catalog):
                lines.append(f'![{figure.caption.replace("[", "(").replace("]", ")")}](evidence/{figure.path.name})\n\n{figure.provenance}' if figure.path else figure.unavailable)
        else:
            lines.append(prefixes[kind] + value)
    (workspace / profile["filename"]).write_text("\n\n".join(lines) + "\n", encoding="utf-8")


def write_general_docx(job, document, workspace):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    profile = PROFILES[profile_id(job)]
    word = Document()
    section = word.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    word.styles["Normal"].font.name = "Arial"
    word.styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2"):
        word.styles[name].font.name = "Arial"
        word.styles[name].font.color.rgb = RGBColor.from_string("571EE6")
    word.core_properties.title = job.get("title") or "Análise"
    word.core_properties.author = "btime Discovery"
    word.core_properties.subject = profile["document"]
    section.header.paragraphs[0].text = "BTIME · DISCOVERY"
    word.add_paragraph(profile["document"].upper(), "Subtitle")
    word.add_heading(job.get("title") or "Análise", 0)
    word.add_paragraph(f'Público: {job.get("audience") or "Equipe"}')
    catalog, seen = figure_catalog(workspace), set()
    for kind, value in document_blocks(job, document, with_figures=True):
        if kind == "figure":
            word_figures(word, value, catalog, seen)
            continue
        if kind in ("h1", "h2"):
            word.add_heading(value, 1 if kind == "h1" else 2)
        else:
            paragraph = word.add_paragraph(value, "List Bullet" if kind == "bullet" else "Normal")
            if kind == "note":
                for run in paragraph.runs:
                    run.italic = True
                    run.font.size = Pt(9)
    word.add_paragraph("Compatibilidade: Word em Arial e PDF em Helvetica; as cores seguem a identidade Btime.")
    word.save(workspace / "documentacao.docx")


def write_general_pdf(job, document, workspace):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    profile = PROFILES[profile_id(job)]
    styles = {
        "p": ParagraphStyle("body", fontName="Helvetica", fontSize=10.5, leading=16, spaceAfter=8),
        "h1": ParagraphStyle("heading", fontName="Helvetica-Bold", fontSize=16, leading=21, spaceBefore=18, spaceAfter=8, keepWithNext=True, textColor=colors.HexColor("#571ee6")),
        "h2": ParagraphStyle("subheading", fontName="Helvetica-Bold", fontSize=12, leading=17, spaceBefore=8, spaceAfter=6, keepWithNext=True),
        "note": ParagraphStyle("note", fontName="Helvetica-Oblique", fontSize=9, leading=13, spaceAfter=8, textColor=colors.HexColor("#62596f")),
        "bullet": ParagraphStyle("bullet", fontSize=10.5, leading=16, leftIndent=12, spaceAfter=6),
    }
    title = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=27, leading=33, spaceAfter=15)
    story = [Paragraph(escape(profile["document"].upper()), styles["note"]),
             Paragraph(escape(job.get("title") or "Análise"), title),
             Paragraph(escape(f'Público: {job.get("audience") or "Equipe"}'), styles["note"]), Spacer(1, 12)]
    catalog, seen = figure_catalog(workspace), set()
    for kind, value in document_blocks(job, document, with_figures=True):
        if kind == "figure":
            story.extend(pdf_figures(value, catalog, seen, max_width=A4[0] - 96))
            continue
        story.append(Paragraph(escape(value).replace("\n", "<br/>"), styles[kind], bulletText="•" if kind == "bullet" else None))
    def decorate(canvas, _):
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(colors.HexColor("#62596f"))
        canvas.drawString(48, A4[1] - 30, "BTIME · DISCOVERY")
        canvas.drawRightString(A4[0] - 48, 28, str(canvas.getPageNumber()))
    path = workspace / "documentacao.pdf"
    SimpleDocTemplate(str(path), pagesize=A4, title=job.get("title") or "Análise", author="btime Discovery", subject=profile["document"],
                      leftMargin=48, rightMargin=48, topMargin=52, bottomMargin=48).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return path


def general_html(job, document, styles, preview, evidence_cards, workspace=None):
    profile = PROFILES[profile_id(job)]
    tags = {"h1": "h2", "h2": "h3", "p": "p", "note": "p", "bullet": "p"}
    catalog = figure_catalog(workspace) if workspace else {}
    content = "".join(html_figures(value, catalog) if kind == "figure" else f'<{tags[kind]} class="{kind}">{escape(value).replace(chr(10), "<br>")}</{tags[kind]}>' for kind, value in document_blocks(job, document, with_figures=True))
    links = "".join(f'<a href="{escape(name)}" download>{escape(label)} · {fmt}</a>' for name, label, fmt in document_catalog(job) if name != "roteiro-cortes.json")
    return (f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">'
            f'<title>{escape(job.get("title") or "Análise")}</title><style>{styles}</style></head><body>'
            f'<header><small>BTIME · DISCOVERY · {escape(profile["document"].upper())}</small><h1>{escape(job.get("title") or "Análise")}</h1>'
            f'<p>Público: {escape(job.get("audience") or "Equipe")}</p><nav class="deliverables">{links}</nav></header>'
            f'<main><section class="panel">{content}</section>{preview}<section><h2>Evidências visuais e falas relacionadas</h2>'
            f'{evidence_cards or "<p>Nenhuma evidência visual disponível. Consulte a transcrição de referência.</p>"}</section></main></body></html>')
