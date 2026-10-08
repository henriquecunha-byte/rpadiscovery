from __future__ import annotations

import base64
import csv
from html import escape
import json
import math
import os
import re
import subprocess
import unicodedata
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openai import OpenAI

from .config import FFMPEG, FFPROBE


VIDEO_EXTENSIONS = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".m4v"}
EXPORT_TYPOGRAPHY_NOTE = "Compatibilidade: Word em Arial e PDF em Helvetica; as cores seguem a identidade Btime."


def find_videos(source: Path) -> list[Path]:
    if source.is_file() and source.suffix.lower() in VIDEO_EXTENSIONS:
        return [source]
    if source.is_dir():
        candidates = sorted((item for item in source.rglob("*") if item.is_file() and item.suffix.lower() in VIDEO_EXTENSIONS), key=lambda item: str(item.relative_to(source)).lower())
        if candidates:
            return candidates
    raise ValueError("Nenhuma gravação de vídeo compatível foi encontrada.")


def prepare_source(source: Path, workspace: Path) -> tuple[Path, list[Path]]:
    videos = find_videos(source)
    if len(videos) == 1:
        return videos[0], videos
    normalized_dir = workspace / "fontes-normalizadas"
    normalized_dir.mkdir(parents=True, exist_ok=True)
    normalized = []
    for index, video in enumerate(videos, 1):
        output = normalized_dir / f"fonte-{index:03d}.mp4"
        duration = probe_duration(video)
        media = probe_media(video)
        if not any(stream.get("codec_type") == "video" for stream in media.get("streams", [])):
            raise ValueError(f"A gravação {video.name} não contém uma faixa de vídeo.")
        has_audio = any(stream.get("codec_type") == "audio" for stream in media.get("streams", []))
        command = [FFMPEG, "-hide_banner", "-loglevel", "error", "-y", "-i", str(video)]
        if not has_audio:
            command += ["-f", "lavfi", "-i", "anullsrc=r=48000:cl=stereo"]
        # Stacks can mix resolutions, frame rates, mono audio and silent clips.
        # The concat demuxer needs identical geometry, stream layout and time base.
        command += [
            "-map", "0:v:0", "-map", "0:a:0" if has_audio else "1:a:0",
            "-vf", "scale=1280:720:force_original_aspect_ratio=decrease:force_divisible_by=2,pad=1280:720:(ow-iw)/2:(oh-ih)/2,setsar=1,fps=30",
            "-c:v", "libx264", "-preset", "veryfast", "-crf", "23", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-b:a", "128k", "-ar", "48000", "-ac", "2", "-af", "apad",
            "-t", f"{duration:.6f}", "-video_track_timescale", "90000", "-movflags", "+faststart", str(output),
        ]
        completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
        if completed.returncode != 0:
            raise RuntimeError(f"Não foi possível preparar a gravação {video.name}: {completed.stderr[-800:]}")
        normalized.append((output, duration))
    concat_list = normalized_dir / "concat.txt"
    # Relative generated names support workspace paths containing apostrophes.
    # Explicit durations avoid accumulating AAC padding between source videos.
    concat_list.write_text("\n".join(f"file '{item.name}'\nduration {duration:.6f}" for item, duration in normalized), encoding="utf-8")
    combined = workspace / "gravacao-consolidada.mp4"
    completed = subprocess.run([
        FFMPEG, "-y", "-f", "concat", "-safe", "0", "-i", str(concat_list),
        "-c", "copy", "-movflags", "+faststart", str(combined),
    ], capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
    if completed.returncode != 0:
        raise RuntimeError(f"Não foi possível consolidar as gravações: {completed.stderr[-800:]}")
    return combined, videos


def probe_duration(video: Path) -> float:
    completed = subprocess.run([FFPROBE, "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1:nk=1", str(video)], capture_output=True, text=True, encoding="utf-8", errors="replace", check=True)
    duration = float(completed.stdout.strip())
    if not math.isfinite(duration) or duration <= 0:
        raise ValueError(f"A gravação {video.name} não possui uma duração válida.")
    return duration


def probe_media(video: Path) -> dict:
    completed = subprocess.run([
        FFPROBE, "-v", "error", "-show_streams", "-show_format", "-of", "json", str(video),
    ], capture_output=True, text=True, encoding="utf-8", errors="replace", check=True)
    return json.loads(completed.stdout)


def extract_evidence(video: Path, output: Path, interval: float = 20.0) -> list[dict]:
    if not math.isfinite(interval) or interval <= 0:
        raise ValueError("O intervalo de captura deve ser maior que zero.")
    duration = probe_duration(video)
    output.mkdir(parents=True, exist_ok=True)
    timestamps = [index * interval for index in range(max(1, int(duration // interval) + 1)) if index * interval < duration]

    def capture(index_and_time: tuple[int, float]) -> dict | None:
        index, timestamp = index_and_time
        filename = f"evidencia-{index:04d}.jpg"
        completed = subprocess.run([
            FFMPEG, "-hide_banner", "-loglevel", "error", "-y", "-ss", f"{timestamp:.3f}",
            "-i", str(video), "-frames:v", "1", "-vf", "scale='min(1280,iw)':-2",
            "-q:v", "3", "-threads", "1", str(output / filename),
        ], capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
        if completed.returncode != 0 or not (output / filename).exists():
            return None
        return {"index": index, "time": round(timestamp, 3), "image": filename}

    workers = max(2, min(6, (os.cpu_count() or 4) // 2))
    evidence = []
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="frame") as executor:
        futures = [executor.submit(capture, item) for item in enumerate(timestamps, 1)]
        for future in as_completed(futures):
            item = future.result()
            if item:
                evidence.append(item)
    return sorted(evidence, key=lambda item: item["index"])


def transcribe(video: Path, workspace: Path) -> list[dict]:
    try:
        from faster_whisper import WhisperModel
        model_name = os.getenv("WHISPER_MODEL", "base")
        model = WhisperModel(
            model_name,
            device="cpu",
            compute_type="int8",
            cpu_threads=max(2, (os.cpu_count() or 4) - 2),
            num_workers=1,
        )
        segments, _ = model.transcribe(
            str(video),
            language="pt",
            vad_filter=True,
            vad_parameters={"min_silence_duration_ms": 500},
            beam_size=1,
            best_of=1,
            condition_on_previous_text=False,
            word_timestamps=False,
        )
        result = [{"start": round(item.start, 3), "end": round(item.end, 3), "text": item.text.strip()} for item in segments]
        (workspace / "transcricao-aviso.txt").unlink(missing_ok=True)
    except Exception as error:
        (workspace / "transcricao-aviso.txt").write_text(f"Transcrição indisponível: {error}", encoding="utf-8")
        result = []
    (workspace / "transcricao.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def nearby_transcript(transcript: list[dict], timestamp: float, radius: float = 18.0) -> str:
    return " ".join(item["text"] for item in transcript if item["end"] >= timestamp - radius and item["start"] <= timestamp + radius)[:1800]


def analyze(evidence: list[dict], transcript: list[dict], workspace: Path, context: str, detail: str) -> list[dict]:
    model = os.getenv("OPENAI_MODEL", "gpt-5.6-luna")
    batches = [evidence[start:start + 10] for start in range(0, len(evidence), 10)]

    def analyze_batch(batch: list[dict]) -> list[dict]:
        content = [{"type": "input_text", "text": (
            "Analise estes frames de uma gravação de discovery de RPA. A fala próxima tem prioridade sobre uma simples mudança visual. "
            "Identifique telas e ações somente quando contribuírem para entender o processo discutido. Ignore troca de câmera, avatar, "
            "rosto de participante, saudação, silêncio e mudança visual sem conteúdo operacional. "
            "Retorne JSON puro no formato {\"steps\":[{\"frame_index\":1,\"title\":\"...\",\"action\":\"...\","
            "\"system\":\"...\",\"field_or_control\":\"...\",\"evidence\":\"o que no frame sustenta a descrição\","
            "\"uncertainty\":\"\"}]}. Não invente cliques, campos ou valores invisíveis. Consolide frames sem mudança relevante. "
            "Não atribua nomes a quem aparece ou fala: esta transcrição não possui identificação confiável de locutor. Use apenas papéis "
            "operacionais quando forem explicitamente informados no contexto do pedido. "
            f"Nível: {detail}. Contexto informado: {context or 'não informado'}."
        )}]
        for item in batch:
            path = workspace / "evidence" / item["image"]
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            content.append({"type": "input_text", "text": f"Frame {item['index']} em {format_time(item['time'])}. Fala próxima: {nearby_transcript(transcript, item['time']) or 'sem fala detectada'}"})
            content.append({"type": "input_image", "image_url": f"data:image/jpeg;base64,{encoded}", "detail": "low"})
        client = OpenAI()
        response = client.responses.create(model=model, input=[{"role": "user", "content": content}])
        text = response.output_text.strip().removeprefix("```json").removesuffix("```").strip()
        if not text.startswith("{"):
            text = text[text.find("{"):text.rfind("}") + 1]
        payload = json.loads(text)
        batch_steps = []
        if not isinstance(payload, dict) or not isinstance(payload.get("steps"), list):
            raise ValueError("A análise visual não retornou uma lista válida de etapas.")
        for step in payload["steps"]:
            if not isinstance(step, dict):
                continue
            frame = next((item for item in batch if str(item["index"]) == str(step.get("frame_index"))), None)
            if frame is None:
                continue
            step["frame_index"] = frame["index"]
            step |= {"time": frame["time"], "timecode": format_time(frame["time"]), "image": frame["image"], "transcript": nearby_transcript(transcript, frame["time"])}
            batch_steps.append(step)
        return batch_steps

    steps = []
    workers = max(1, min(3, len(batches)))
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="vision") as executor:
        futures = [executor.submit(analyze_batch, batch) for batch in batches]
        for future in as_completed(futures):
            steps.extend(future.result())
    return sorted(steps, key=lambda item: float(item.get("time", 0)))


def fallback_documentation(job: dict, steps: list[dict], warning: str = "") -> dict:
    systems = sorted({str(step.get("system") or "").strip() for step in steps if str(step.get("system") or "").strip()})
    flow = []
    for sequence, step in enumerate(steps, 1):
        flow.append({
            "sequence": sequence,
            "title": step.get("title") or f"Etapa {sequence}",
            "description": step.get("action") or "Ação não descrita.",
            "actor": "Não identificado",
            "system": step.get("system") or "Não identificado",
            "input": "Não identificado",
            "output": "Não identificado",
            "decision_or_exception": "",
            "evidence_refs": [step.get("frame_index")] if step.get("frame_index") else [],
        })
    limitations = ["A documentação foi montada somente com os fatos observáveis na gravação."]
    if warning:
        limitations.append(f"A síntese automática não pôde ser concluída: {warning}")
    return {
        "executive_summary": "Registro preliminar das etapas observadas. A síntese do processo e seus requisitos precisam de validação antes do uso operacional.",
        "objective": "Objetivo de negócio não confirmado. Validar com o responsável pelo processo.",
        "scope": {"in_scope": [], "out_of_scope": []},
        "actors": [],
        "systems": [{"name": name, "purpose": "Finalidade não explicitada."} for name in systems],
        "prerequisites": [],
        "inputs": [],
        "outputs": [],
        "business_rules": [],
        "process_flow": flow,
        "exceptions": [],
        "risks_and_controls": [],
        "open_questions": ["Validar com o responsável pelo processo os pontos que não foram explicitados na gravação."],
        "automation_opportunities": [],
        "preview_moments": [],
        "limitations": limitations,
    }


def normalize_documentation(document, job: dict, steps: list[dict]) -> dict:
    baseline = fallback_documentation(job, steps)
    if not isinstance(document, dict):
        return baseline
    for key in ("executive_summary", "objective"):
        if not isinstance(document.get(key), str) or not document[key].strip():
            document[key] = baseline[key]
    scope = document.get("scope")
    if not isinstance(scope, dict):
        scope = baseline["scope"]
    for key in ("in_scope", "out_of_scope"):
        if not isinstance(scope.get(key), list):
            scope[key] = []
    document["scope"] = scope
    list_fields = (
        "actors", "systems", "prerequisites", "inputs", "outputs", "business_rules",
        "process_flow", "exceptions", "risks_and_controls", "open_questions",
        "automation_opportunities", "limitations",
        "preview_moments",
    )
    for key in list_fields:
        if not isinstance(document.get(key), list):
            document[key] = baseline[key]
    # Model JSON can be valid JSON but still violate the document schema. Keep
    # usable facts and prevent a malformed list item from breaking every export.
    object_fields = ("actors", "systems", "inputs", "outputs", "business_rules", "process_flow",
                     "exceptions", "risks_and_controls", "automation_opportunities", "preview_moments")
    for key in object_fields:
        document[key] = [dict(item) for item in document[key] if isinstance(item, dict) and item]
    for key in ("prerequisites", "open_questions", "limitations"):
        document[key] = [item.strip() for item in document[key] if isinstance(item, str) and item.strip()]
    for key in ("in_scope", "out_of_scope"):
        scope[key] = [item.strip() for item in scope[key] if isinstance(item, str) and item.strip()]
    known_refs = {str(item["frame_index"]) for item in steps if item.get("frame_index") is not None}
    for key in object_fields:
        for item in document[key]:
            for field, value in list(item.items()):
                if field == "evidence_refs":
                    item[field] = [ref for ref in value if str(ref) in known_refs] if isinstance(value, list) else []
                elif isinstance(value, (dict, list)):
                    item[field] = ""
            if key in ("business_rules", "process_flow"):
                item.setdefault("evidence_refs", [])
    if not document["process_flow"]:
        document["process_flow"] = baseline["process_flow"]
    return document


def synthesize_documentation(job: dict, steps: list[dict], transcript: list[dict], workspace: Path) -> dict:
    transcript_lines = "\n".join(f"[{format_time(item['start'])}–{format_time(item['end'])}] {item['text']}" for item in transcript)
    observed_steps = [{
        "frame_index": item.get("frame_index"),
        "timecode": item.get("timecode"),
        "title": item.get("title"),
        "action": item.get("action"),
        "system": item.get("system"),
        "field_or_control": item.get("field_or_control"),
        "evidence": item.get("evidence"),
        "uncertainty": item.get("uncertainty"),
    } for item in steps]
    prompt = f'''Você é analista de processos e discovery de RPA. Transforme os insumos abaixo em documentação operacional estruturada, e não em resumo de transcrição.

REGRAS DE CONFIABILIDADE
- Use somente fatos demonstrados visualmente, explicitamente falados ou informados no contexto do pedido.
- Não invente regras, responsáveis, integrações, frequências, volumes, credenciais, exceções ou sistemas.
- O áudio não possui diarização confiável. Não atribua falas ou ações a uma pessoa pelo nome, mesmo que um nome apareça na interface ou seja citado. Use papéis como "área solicitante", "analista" ou "responsável pelo processo" somente quando o papel estiver explícito. Nomes fornecidos diretamente no contexto do pedido podem ser usados.
- Quando algo necessário não estiver claro, registre em open_questions ou use "Não identificado".
- Consolide ações de tela em etapas de negócio compreensíveis. A transcrição e os frames são evidências, não a estrutura principal.
- evidence_refs deve conter somente frame_index existentes nos passos observados.
- Diferencie o processo atual de oportunidades futuras de automação.

REGRAS PARA O PREVIEW COM CORTES
- O objetivo NÃO é resumir a reunião, criar destaques ou atingir uma duração curta. O objetivo é separar o que pertence ao assunto pedido do que pertence a outros assuntos.
- Faça uma decisão contextual binária para cada bloco da conversa: "é sobre o assunto definido no Contexto informado?" Se sim, mantenha integralmente. Se não, remova.
- Considere relevante não apenas a menção literal ao nome do cliente ou projeto, mas também perguntas, respostas, contrapontos, exemplos, demonstrações de tela, decisões, próximos passos e explicações que dependam desse contexto.
- Preserve o início que apresenta uma dúvida e o final que conclui a resposta. Não retire falas intermediárias só porque não repetem o nome do assunto.
- Quando a conversa permanecer no mesmo assunto, gere um bloco contínuo e amplo. Não fragmente uma discussão relevante em pequenos melhores momentos.
- O frame serve para confirmar ou ilustrar o que está sendo falado; sozinho, não deve decidir o corte.
- Exclua apenas blocos claramente dedicados a outro assunto, além de saudações, espera, silêncio, problemas técnicos sem conteúdo e conversas paralelas sem relação com o pedido.
- Não use quantidade de momentos nem percentual da duração como meta. Se grande parte da reunião for sobre o assunto pedido, grande parte deve permanecer.
- Cada momento deve começar pouco antes da primeira fala relevante e terminar depois que a última ideia relevante for concluída, sem cortar no meio da fala.
- start e end devem usar timecodes existentes na transcrição e end deve ser posterior a start.

PEDIDO
Título: {job.get('title', '')}
Contexto informado: {job.get('process_context') or 'Não informado'}
Público: {job.get('audience') or 'Equipe de RPA'}
Nível: {job.get('detail_level') or 'operacional'}

PASSOS OBSERVADOS
{json.dumps(observed_steps, ensure_ascii=False)}

TRANSCRIÇÃO COM TIMECODE
{transcript_lines or 'Transcrição indisponível.'}

Retorne JSON puro exatamente com esta estrutura:
{{
  "executive_summary": "síntese do processo em linguagem de negócio",
  "objective": "resultado que o processo busca produzir",
  "scope": {{"in_scope": ["..."], "out_of_scope": ["..."]}},
  "actors": [{{"name": "...", "responsibility": "..."}}],
  "systems": [{{"name": "...", "purpose": "..."}}],
  "prerequisites": ["..."],
  "inputs": [{{"name": "...", "source": "...", "required": "..."}}],
  "outputs": [{{"name": "...", "destination": "..."}}],
  "business_rules": [{{"rule": "...", "evidence_refs": [1]}}],
  "process_flow": [{{"sequence": 1, "title": "...", "description": "...", "actor": "...", "system": "...", "input": "...", "output": "...", "decision_or_exception": "...", "evidence_refs": [1]}}],
  "exceptions": [{{"scenario": "...", "handling": "...", "status": "confirmado ou a validar"}}],
  "risks_and_controls": [{{"risk": "...", "impact": "...", "control": "..."}}],
  "open_questions": ["..."],
  "automation_opportunities": [{{"opportunity": "...", "benefit": "...", "dependency": "..."}}],
  "preview_moments": [{{"start": "00:10:05", "end": "00:11:20", "title": "explicação do trecho", "reason": "por que esse conteúdo falado é necessário"}}],
  "limitations": ["..."]
}}'''
    try:
        client = OpenAI()
        response = client.responses.create(model=os.getenv("OPENAI_MODEL", "gpt-5.6-luna"), input=prompt)
        text = response.output_text.strip().removeprefix("```json").removesuffix("```").strip()
        if not text.startswith("{"):
            text = text[text.find("{"):text.rfind("}") + 1]
        document = json.loads(text)
        baseline = fallback_documentation(job, steps)
        for key, value in baseline.items():
            document.setdefault(key, value)
        (workspace / "documentacao-aviso.txt").unlink(missing_ok=True)
        return normalize_documentation(document, job, steps)
    except Exception as error:
        (workspace / "documentacao-aviso.txt").write_text(f"Síntese estruturada indisponível: {error}", encoding="utf-8")
        return fallback_documentation(job, steps, str(error))


def fallback_process_prompt(job: dict) -> str:
    title = str(job.get("title") or "Processo operacional").strip()
    audience = str(job.get("audience") or "Equipe de RPA").strip()
    detail = str(job.get("detail_level") or "operacional").strip()
    hint = str(job.get("process_context") or "").strip()
    sources = [str(item).strip() for item in job.get("sources") or [] if str(item).strip()]
    subject = hint or f"o processo tratado em “{title}”"
    lines = [
        f"Assunto do discovery: {subject}.",
        f"Documente esse processo para {audience}, em nível {detail}, a partir do que for demonstrado em tela e falado na gravação.",
    ]
    if sources:
        listed = ", ".join(sources[:6]) + (" e demais arquivos da pilha" if len(sources) > 6 else "")
        lines.append(f"Gravações enviadas: {listed}.")
    lines += [
        "Mantenha no preview todo bloco falado que pertença a esse assunto, da pergunta inicial até a conclusão da resposta, incluindo contrapontos, exemplos, demonstrações de tela, decisões e próximos passos.",
        "Descarte apenas os blocos claramente dedicados a outro assunto, além de saudação, espera, silêncio, problema técnico sem conteúdo e conversa paralela.",
        "Registre sistemas, telas, campos, entradas, saídas, regras, exceções e riscos somente quando estiverem visíveis na tela ou ditos em voz alta; o que ficar implícito deve virar ponto a validar.",
        "Não atribua falas ou ações a pessoas pelo nome: esta gravação não possui identificação confiável de locutor.",
    ]
    return "\n".join(lines)


def suggest_process_prompt(job: dict) -> dict:
    sources = [str(item).strip() for item in job.get("sources") or [] if str(item).strip()]
    fallback = fallback_process_prompt(job)
    instruction = f'''Você escreve o campo "Contexto informado pela equipe" de uma ferramenta de documentação pós-discovery de RPA. Produza o texto que a equipe usaria para orientar a análise da gravação que será enviada.

COMO A FERRAMENTA USA ESSE TEXTO
- Ele acompanha cada lote de frames na análise visual e define o que conta como tela ou ação relevante do processo.
- Ele guia a síntese estruturada: objetivo, escopo, atores, sistemas, pré-requisitos, entradas, saídas, regras de negócio, fluxo operacional, exceções, riscos, oportunidades de automação e pontos a validar.
- Ele decide o corte do preview em vídeo: cada bloco falado é mantido ou removido conforme pertencer ou não ao assunto declarado aqui. Não existe meta de duração.
- A transcrição é local e não possui diarização confiável, então o texto não pode pedir atribuição de falas por nome.

REGRAS DO TEXTO
- Português do Brasil, texto corrido em linhas curtas, sem título, sem markdown, sem numeração e sem aspas ao redor de tudo.
- Entre 6 e 10 linhas, cada uma uma orientação objetiva e verificável.
- A primeira linha declara o assunto do discovery de forma inequívoca, porque é ela que separa o que fica e o que sai do preview.
- Inclua o que deve ser mantido integralmente, o que deve ser descartado, quais elementos operacionais procurar e como tratar o que não estiver explícito.
- Não invente cliente, sistema, número, prazo, regra ou responsável que não tenha sido informado abaixo. Quando faltar informação, escreva a orientação de forma genérica em vez de preencher com suposição.

PEDIDO DA EQUIPE
Título do trabalho: {job.get('title') or 'Não informado'}
Público: {job.get('audience') or 'Equipe de RPA'}
Nível de detalhe: {job.get('detail_level') or 'operacional'}
Rascunho ou assunto já escrito pela equipe: {job.get('process_context') or 'Nada foi escrito ainda'}
Arquivos na pilha: {', '.join(sources) if sources else 'Nenhum arquivo selecionado ainda'}

Responda somente com o texto do contexto.'''
    try:
        client = OpenAI()
        response = client.responses.create(model=os.getenv("OPENAI_MODEL", "gpt-5.6-luna"), input=instruction)
        text = response.output_text.strip().removeprefix("```").removesuffix("```").strip()
        if len(text) < 40:
            raise ValueError("A sugestão retornada ficou curta demais para orientar a análise.")
        return {"prompt": text[:8000], "generated": True, "warning": ""}
    except Exception as error:
        return {"prompt": fallback, "generated": False, "warning": str(error)}


def format_time(seconds: float) -> str:
    value = max(0, round(seconds))
    return f"{value // 3600:02d}:{(value % 3600) // 60:02d}:{value % 60:02d}"


def parse_timecode(value) -> float:
    if isinstance(value, bool):
        raise ValueError(f"Timecode inválido: {value}")
    if isinstance(value, (int, float)):
        seconds = float(value)
        if not math.isfinite(seconds) or seconds < 0:
            raise ValueError(f"Timecode inválido: {value}")
        return seconds
    parts = str(value or "").strip().split(":")
    if len(parts) > 3 or any(not part.replace(".", "", 1).isdigit() for part in parts):
        raise ValueError(f"Timecode inválido: {value}")
    if len(parts) > 1 and (any(not part.isdigit() for part in parts[:-1]) or float(parts[-1]) >= 60):
        raise ValueError(f"Timecode inválido: {value}")
    if len(parts) == 3 and int(parts[1]) >= 60:
        raise ValueError(f"Timecode inválido: {value}")
    seconds = 0.0
    for part in parts:
        seconds = seconds * 60 + float(part)
    return seconds


def preview_ranges(steps: list[dict], duration: float, before: float = 6.0, after: float = 12.0) -> list[tuple[float, float]]:
    ranges = []
    for step in sorted(steps, key=lambda item: float(item.get("time", 0))):
        start = max(0.0, float(step.get("time", 0)) - before)
        end = min(duration, float(step.get("time", 0)) + after)
        if end <= start:
            continue
        if ranges and start <= ranges[-1][1] + 5.0:
            ranges[-1] = (ranges[-1][0], max(ranges[-1][1], end))
        else:
            ranges.append((start, end))
    return ranges


def preview_ranges_from_moments(moments: list[dict], duration: float, before: float = 4.0, after: float = 6.0, merge_gap: float = 0.0) -> list[tuple[float, float]]:
    ranges = []
    for moment in moments or []:
        if not isinstance(moment, dict):
            continue
        try:
            raw_start = parse_timecode(moment.get("start"))
            raw_end = parse_timecode(moment.get("end"))
        except (TypeError, ValueError):
            continue
        if raw_end <= raw_start or raw_start >= duration:
            continue
        start = max(0.0, raw_start - before)
        end = min(duration, raw_end + after)
        ranges.append((start, end))
    merged = []
    for start, end in sorted(ranges):
        if merged and start <= merged[-1][1] + merge_gap:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return merged


def build_preview_plan(video: Path, steps: list[dict], moments: list[dict] | None = None) -> dict:
    duration = probe_duration(video)
    spoken_ranges = preview_ranges_from_moments(moments or [], duration)
    # An explicit empty contextual selection means no relevant speech, not
    # permission to include unrelated screen changes. None is legacy-only.
    ranges = spoken_ranges if moments is not None else preview_ranges(steps, duration)
    return {
        "created": bool(ranges),
        "duration": round(sum(end - start for start, end in ranges), 3),
        "ranges": [{"start": round(start, 3), "end": round(end, 3)} for start, end in ranges],
        "selection_basis": "contextual_spoken_content" if moments is not None else ("visual_evidence_fallback" if ranges else "none"),
        "moments": moments or [],
    }


def render_preview(video: Path, ranges: list[tuple[float, float]], workspace: Path, output_path: Path, plan_path: Path, selection_basis: str, moments: list[dict] | None = None) -> dict:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    plan_path.parent.mkdir(parents=True, exist_ok=True)
    if not ranges:
        output_path.unlink(missing_ok=True)
        result = {"created": False, "duration": 0, "ranges": [], "selection_basis": selection_basis, "moments": moments or [], "file": None}
        plan_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return result
    segments_dir = workspace / "preview-segments" / output_path.stem
    segments_dir.mkdir(parents=True, exist_ok=True)
    segments = []
    for index, (start, end) in enumerate(ranges, 1):
        segment = segments_dir / f"segment-{index:03d}.mp4"
        completed = subprocess.run([
            FFMPEG, "-y", "-ss", f"{start:.3f}", "-i", str(video), "-t", f"{end - start:.3f}",
            "-map", "0:v:0", "-map", "0:a:0?", "-vf", "scale='min(1280,trunc(iw/2)*2)':-2",
            "-c:v", "libx264", "-preset", "ultrafast", "-crf", "27", "-pix_fmt", "yuv420p", "-c:a", "aac",
            "-b:a", "96k", "-movflags", "+faststart", str(segment),
        ], capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
        if completed.returncode != 0:
            raise RuntimeError(f"Não foi possível montar o trecho {index} do preview: {completed.stderr[-800:]}")
        segments.append(segment)
    concat_list = segments_dir / "concat.txt"
    concat_list.write_text("\n".join(f"file '{item.name}'\nduration {end - start:.6f}" for item, (start, end) in zip(segments, ranges)), encoding="utf-8")
    completed = subprocess.run([
        FFMPEG, "-y", "-f", "concat", "-safe", "0", "-i", str(concat_list),
        "-c", "copy", "-movflags", "+faststart", str(output_path),
    ], capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
    if completed.returncode != 0:
        raise RuntimeError(f"Não foi possível finalizar o preview: {completed.stderr[-800:]}")
    result = {
        "created": True,
        "duration": round(sum(end - start for start, end in ranges), 3),
        "ranges": [{"start": round(start, 3), "end": round(end, 3)} for start, end in ranges],
        "selection_basis": selection_basis,
        "moments": moments or [],
        "file": output_path.relative_to(workspace).as_posix(),
    }
    plan_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def create_preview(video: Path, steps: list[dict], workspace: Path, moments: list[dict] | None = None) -> dict:
    plan = build_preview_plan(video, steps, moments)
    ranges = [(item["start"], item["end"]) for item in plan["ranges"]]
    return render_preview(video, ranges, workspace, workspace / "preview-processo.mp4", workspace / "roteiro-cortes.json", plan["selection_basis"], moments)


def preview_slug(video: Path, index: int) -> str:
    normalized = unicodedata.normalize("NFKD", video.stem).encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "-", normalized).strip("-").lower()[:70] or "video"
    return f"{index:03d}-{cleaned}"


def create_job_previews(analysis_video: Path, source_videos: list[Path], steps: list[dict], workspace: Path, moments: list[dict] | None = None) -> tuple[dict, list[dict]]:
    if len(source_videos) == 1:
        preview = create_preview(analysis_video, steps, workspace, moments)
        source_preview = dict(preview)
        source_preview |= {"index": 1, "source_name": source_videos[0].name, "source_path": str(source_videos[0]), "global_offset": 0.0}
        (workspace / "roteiro-cortes.json").write_text(json.dumps(source_preview, ensure_ascii=False, indent=2), encoding="utf-8")
        return preview, [source_preview]

    plan = build_preview_plan(analysis_video, steps, moments)
    global_ranges = [(item["start"], item["end"]) for item in plan["ranges"]]
    source_previews = []
    offset = 0.0
    for index, source in enumerate(source_videos, 1):
        source_duration = probe_duration(source)
        source_end = offset + source_duration
        local_ranges = []
        for start, end in global_ranges:
            local_start = max(start, offset)
            local_end = min(end, source_end)
            if local_end - local_start > 0.05:
                local_ranges.append((local_start - offset, local_end - offset))
        slug = preview_slug(source, index)
        source_preview = render_preview(
            source,
            local_ranges,
            workspace,
            workspace / "previews" / f"{slug}-preview.mp4",
            workspace / "roteiros" / f"{slug}-cortes.json",
            plan["selection_basis"],
            moments,
        )
        source_preview |= {
            "index": index,
            "source_name": source.name,
            "source_path": str(source),
            "source_duration": round(source_duration, 3),
            "global_offset": round(offset, 3),
            "moments_timeline": "consolidated",
        }
        (workspace / "roteiros" / f"{slug}-cortes.json").write_text(json.dumps(source_preview, ensure_ascii=False, indent=2), encoding="utf-8")
        source_previews.append(source_preview)
        offset = source_end

    (workspace / "preview-processo.mp4").unlink(missing_ok=True)
    plan["created"] = any(item["created"] for item in source_previews)
    plan["previews"] = source_previews
    (workspace / "roteiro-cortes.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    return plan, source_previews


def create_package(workspace: Path):
    package = workspace / "entrega-completa.zip"
    temporary = workspace / "entrega-completa.zip.tmp"
    included = [
        "relatorio.html", "relatorio.json", "documentacao-processo.json",
        "documentacao-processo.docx", "documentacao-processo.pdf", "procedimento-operacional.md", "requisitos-rpa.md", "matriz-evidencias.csv",
        "transcricao.json", "preview-processo.mp4", "roteiro-cortes.json", "transcricao-aviso.txt", "documentacao-aviso.txt", "documentacao-pdf-aviso.txt",
    ]
    current_files = None
    plan_path = workspace / "roteiro-cortes.json"
    if plan_path.exists():
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        current_files = set()
        previews = plan.get("previews", [plan])
        for item in previews:
            if item.get("created") and item.get("file"):
                current_files.add(item["file"])
            if item.get("index") and item.get("source_name"):
                current_files.add(f"roteiros/{preview_slug(Path(item['source_name']), item['index'])}-cortes.json")
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for name in included:
            path = workspace / name
            if name == "preview-processo.mp4" and current_files is not None and name not in current_files:
                continue
            if path.exists():
                compression = zipfile.ZIP_STORED if path.suffix.lower() == ".mp4" else zipfile.ZIP_DEFLATED
                archive.write(path, arcname=name, compress_type=compression)
        for path in sorted((workspace / "evidence").glob("*.jpg")):
            archive.write(path, arcname=f"evidence/{path.name}", compress_type=zipfile.ZIP_STORED)
        for folder in ("previews", "roteiros"):
            for path in sorted((workspace / folder).glob("*")):
                if path.is_file() and (current_files is None or f"{folder}/{path.name}" in current_files):
                    compression = zipfile.ZIP_STORED if path.suffix.lower() == ".mp4" else zipfile.ZIP_DEFLATED
                    archive.write(path, arcname=f"{folder}/{path.name}", compress_type=compression)
    temporary.replace(package)


def markdown_value(value) -> str:
    return str(value or "Não identificado").replace("|", "\\|").replace("\n", " ")


def write_docx_document(job: dict, document: dict, workspace: Path):
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor, Twips
    except ImportError as error:
        (workspace / "documentacao-aviso.txt").write_text(f"Documento Word indisponível: {error}", encoding="utf-8")
        return

    navy = RGBColor.from_string("060315")
    teal = RGBColor.from_string("571EE6")
    muted = RGBColor.from_string("62596F")
    word = Document()
    word.core_properties.title = str(job.get("title") or "Processo")
    word.core_properties.subject = "Documentação pós-discovery de RPA"
    word.core_properties.author = "Btime RPA Docs"
    section = word.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    def set_font(run, size=None, color=None, bold=None, italic=None):
        run.font.name = "Arial"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic

    normal = word.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.color.rgb = navy
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, navy, 18, 10), ("Heading 2", 13, navy, 14, 7), ("Heading 3", 12, muted, 10, 5),
    ):
        style = word.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "BTIME · DOCUMENTAÇÃO PÓS-DISCOVERY DE RPA"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.runs[0], 8.5, muted, True)
    footer = section.footer.paragraphs[0]
    footer.text = str(job.get("title") or "Documento de processo")
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.runs[0], 8.5, muted)
    set_font(footer.add_run(" · página "), 8.5, muted)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    kicker = word.add_paragraph()
    kicker.paragraph_format.space_before = Pt(48)
    kicker.paragraph_format.space_after = Pt(8)
    set_font(kicker.add_run("DOCUMENTO DE PROCESSO E REQUISITOS PARA RPA"), 9, teal, True)
    title = word.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run(str(job.get("title") or "Processo")), 28, navy, True)
    subtitle = word.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    set_font(subtitle.add_run(f"Público: {job.get('audience') or 'Equipe de RPA'}  |  Nível: {job.get('detail_level') or 'operacional'}"), 10.5, muted)
    lead = word.add_paragraph()
    lead.paragraph_format.space_after = Pt(18)
    set_font(lead.add_run(str(document.get("executive_summary") or "Não identificado na gravação.")), 13, navy)
    note = word.add_paragraph()
    set_font(note.add_run("Este documento consolida o processo demonstrado. A transcrição e os prints são evidências de apoio e permanecem separados no pacote."), 9.5, muted, italic=True)
    if job.get("process_context"):
        guide_heading = word.add_heading("Guia recebido para esta análise", level=2)
        guide_heading.paragraph_format.space_before = Pt(20)
        guide = word.add_paragraph(str(job["process_context"]))
        guide.paragraph_format.space_after = Pt(8)
        guidance_note = word.add_paragraph()
        set_font(guidance_note.add_run("Orientação da equipe para documentação e cortes; não equivale a um objetivo de negócio confirmado."), 9, muted, italic=True)
    word.add_page_break()

    def heading(text: str, level: int = 1):
        return word.add_heading(text, level=level)

    def paragraph(text, bold_label: str = ""):
        item = word.add_paragraph()
        if bold_label:
            set_font(item.add_run(bold_label), 11, navy, True)
        item.add_run(str(text or "Não identificado na gravação."))
        return item

    def bullet_items(items: list, empty: str = "Não identificado na gravação."):
        values = items or [empty]
        for value in values:
            item = word.add_paragraph(style="List Bullet")
            item.paragraph_format.left_indent = Inches(0.375)
            item.paragraph_format.first_line_indent = Inches(-0.188)
            item.paragraph_format.space_after = Pt(4)
            item.paragraph_format.line_spacing = 1.25
            item.add_run(str(value))

    def shade_cell(cell, fill: str):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shading)

    def ensure_child(parent, tag: str):
        child = parent.find(qn(tag))
        if child is None:
            child = OxmlElement(tag)
            parent.append(child)
        return child

    def apply_table_geometry(table, widths: list[float]):
        widths_dxa = [int(round(width * 1440)) for width in widths]
        widths_dxa[-1] += 9360 - sum(widths_dxa)
        properties = table._tbl.tblPr
        table_width = ensure_child(properties, "w:tblW")
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), "9360")
        table_indent = ensure_child(properties, "w:tblInd")
        table_indent.set(qn("w:type"), "dxa")
        table_indent.set(qn("w:w"), "120")
        layout = ensure_child(properties, "w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths_dxa:
            grid_column = OxmlElement("w:gridCol")
            grid_column.set(qn("w:w"), str(width))
            grid.append(grid_column)
        for column_index, width in enumerate(widths_dxa):
            table.columns[column_index].width = Twips(width)
            for row in table.rows:
                cell = row.cells[column_index]
                cell.width = Twips(width)
                cell_width = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcW")
                cell_width.set(qn("w:type"), "dxa")
                cell_width.set(qn("w:w"), str(width))
                margins = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcMar")
                for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                    margin = ensure_child(margins, f"w:{side}")
                    margin.set(qn("w:type"), "dxa")
                    margin.set(qn("w:w"), str(value))

    def data_table(headers: list[str], rows: list[list], widths: list[float]):
        table = word.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        row_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        row_properties.append(repeat_header)
        for index, (header_text, width) in enumerate(zip(headers, widths)):
            cell = table.rows[0].cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "F0EAFF")
            set_font(cell.paragraphs[0].add_run(header_text), 9.5, navy, True)
        for values in rows:
            cells = table.add_row().cells
            for index, (value, width) in enumerate(zip(values, widths)):
                cells[index].width = Inches(width)
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_font(cells[index].paragraphs[0].add_run(str(value or "Não identificado")), 9.5, None)
        apply_table_geometry(table, widths)
        word.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    heading("1. Objetivo")
    paragraph(document.get("objective"))
    heading("2. Escopo")
    heading("Incluído", 2)
    bullet_items(document.get("scope", {}).get("in_scope", []))
    heading("Fora do escopo", 2)
    bullet_items(document.get("scope", {}).get("out_of_scope", []))
    heading("3. Atores e responsabilidades")
    actor_rows = [[item.get("name"), item.get("responsibility")] for item in document.get("actors", [])] or [["Não identificado", "Validar com a área responsável"]]
    data_table(["Ator", "Responsabilidade"], actor_rows, [1.7, 4.8])
    heading("4. Sistemas envolvidos")
    system_rows = [[item.get("name"), item.get("purpose")] for item in document.get("systems", [])] or [["Não identificado", "Validar durante o refinamento"]]
    data_table(["Sistema", "Finalidade no processo"], system_rows, [1.7, 4.8])
    heading("5. Pré-requisitos")
    bullet_items(document.get("prerequisites", []))
    heading("6. Entradas e saídas")
    heading("Entradas", 2)
    input_rows = [[item.get("name"), item.get("source"), item.get("required")] for item in document.get("inputs", [])] or [["Não identificado", "A validar", "A validar"]]
    data_table(["Entrada", "Origem", "Obrigatoriedade"], input_rows, [2.5, 2.4, 1.6])
    heading("Saídas", 2)
    output_rows = [[item.get("name"), item.get("destination")] for item in document.get("outputs", [])] or [["Não identificado", "A validar"]]
    data_table(["Saída", "Destino"], output_rows, [3.25, 3.25])
    heading("7. Fluxo operacional")
    for index, item in enumerate(document.get("process_flow", []), 1):
        heading(f"{item.get('sequence') or index}. {item.get('title') or 'Etapa'}", 2)
        paragraph(item.get("description"))
        paragraph(item.get("actor"), "Responsável: ")
        paragraph(item.get("system"), "Sistema: ")
        paragraph(item.get("input"), "Entrada: ")
        paragraph(item.get("output"), "Saída: ")
        if item.get("decision_or_exception"):
            paragraph(item.get("decision_or_exception"), "Decisão ou exceção: ")
        refs = ", ".join(str(ref) for ref in item.get("evidence_refs", [])) or "sem referência visual direta"
        evidence = word.add_paragraph()
        set_font(evidence.add_run(f"Evidências: {refs}"), 9, muted, italic=True)
    heading("8. Regras de negócio")
    bullet_items([item.get("rule") for item in document.get("business_rules", []) if item.get("rule")], "Nenhuma regra foi explicitada na gravação.")
    heading("9. Exceções e tratamentos")
    for item in document.get("exceptions", []):
        paragraph(item.get("scenario"), "Cenário: ")
        paragraph(item.get("handling"), "Tratamento: ")
        paragraph(item.get("status"), "Status: ")
    if not document.get("exceptions"):
        paragraph("Nenhuma exceção foi explicitada na gravação.")
    heading("10. Riscos e controles")
    for item in document.get("risks_and_controls", []):
        paragraph(item.get("risk"), "Risco: ")
        paragraph(item.get("impact"), "Impacto: ")
        paragraph(item.get("control"), "Controle: ")
    if not document.get("risks_and_controls"):
        paragraph("Nenhum risco ou controle foi explicitado na gravação.")
    heading("11. Oportunidades de automação")
    for item in document.get("automation_opportunities", []):
        heading(item.get("opportunity") or "Oportunidade", 2)
        paragraph(item.get("benefit"), "Benefício: ")
        paragraph(item.get("dependency"), "Dependência: ")
    if not document.get("automation_opportunities"):
        paragraph("Nenhuma oportunidade foi confirmada com os insumos disponíveis.")
    heading("12. Pontos a validar")
    bullet_items(document.get("open_questions", []), "Nenhum ponto adicional registrado.")
    heading("13. Limitações da análise")
    bullet_items(document.get("limitations", []))
    set_font(word.add_paragraph().add_run(EXPORT_TYPOGRAPHY_NOTE), 8.5, muted)
    word.save(workspace / "documentacao-processo.docx")


def write_pdf_document(job: dict, document: dict, workspace: Path) -> Path | None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as error:
        (workspace / "documentacao-pdf-aviso.txt").write_text(f"Exportação em PDF indisponível: {error}", encoding="utf-8")
        return None

    navy = colors.HexColor("#060315")
    teal = colors.HexColor("#571ee6")
    muted = colors.HexColor("#62596f")
    ink = navy
    path = workspace / "documentacao-processo.pdf"
    title_text = str(job.get("title") or "Processo")

    base = ParagraphStyle("Corpo", fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=5, textColor=ink)
    styles = {
        "body": base,
        "kicker": ParagraphStyle("Kicker", parent=base, fontName="Helvetica-Bold", fontSize=9, textColor=teal, spaceAfter=8),
        "cover": ParagraphStyle("Capa", parent=base, fontName="Helvetica-Bold", fontSize=27, leading=31, textColor=navy, spaceAfter=10),
        "subtitle": ParagraphStyle("Subtitulo", parent=base, fontSize=10.5, textColor=muted, spaceAfter=22),
        "lead": ParagraphStyle("Abertura", parent=base, fontSize=13, leading=18, textColor=navy, spaceAfter=18),
        "note": ParagraphStyle("Nota", parent=base, fontName="Helvetica-Oblique", fontSize=9.5, textColor=muted),
        "h1": ParagraphStyle("Titulo1", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=navy, spaceBefore=14, spaceAfter=7, keepWithNext=1),
        "h2": ParagraphStyle("Titulo2", parent=base, fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=navy, spaceBefore=10, spaceAfter=5, keepWithNext=1),
        "bullet": ParagraphStyle("Marcador", parent=base, leftIndent=16, bulletIndent=4, spaceAfter=4),
        "evidence": ParagraphStyle("Evidencia", parent=base, fontName="Helvetica-Oblique", fontSize=9, textColor=muted, spaceAfter=10),
        "th": ParagraphStyle("Cabecalho", parent=base, fontName="Helvetica-Bold", fontSize=9.5, leading=13, textColor=navy, spaceAfter=0),
        "td": ParagraphStyle("Celula", parent=base, fontSize=9.5, leading=13, spaceAfter=0),
    }

    def text(value, style="body", label=""):
        content = escape(str(value or "Não identificado na gravação."))
        if label:
            content = f"<b>{escape(label)}</b> {content}"
        return Paragraph(content, styles[style])

    def bullets(items: list, empty: str = "Não identificado na gravação."):
        return [Paragraph(escape(str(item)), styles["bullet"], bulletText="•") for item in (items or [empty])]

    def data_table(headers: list[str], rows: list[list], widths: list[float]):
        data = [[Paragraph(escape(header), styles["th"]) for header in headers]]
        data += [[Paragraph(escape(str(value or "Não identificado")), styles["td"]) for value in row] for row in rows]
        table = Table(data, colWidths=[width * inch for width in widths], repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0eaff")),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#e1ddeb")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ]))
        return [table, Spacer(1, 8)]

    def decorate(canvas, document_template):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8.5)
        canvas.setFillColor(muted)
        canvas.drawString(inch, LETTER[1] - 0.72 * inch, "BTIME · DOCUMENTAÇÃO PÓS-DISCOVERY DE RPA")
        canvas.setFont("Helvetica", 8.5)
        footer_title = title_text
        suffix = f" · página {canvas.getPageNumber()}"
        while footer_title and canvas.stringWidth(footer_title + suffix, "Helvetica", 8.5) > LETTER[0] - 2 * inch:
            footer_title = footer_title[:-2].rstrip() + "…" if len(footer_title) > 2 else ""
        canvas.drawRightString(LETTER[0] - inch, 0.62 * inch, footer_title + suffix)
        canvas.setStrokeColor(colors.HexColor("#e1ddeb"))
        canvas.line(inch, LETTER[1] - 0.82 * inch, LETTER[0] - inch, LETTER[1] - 0.82 * inch)
        canvas.restoreState()

    story = [
        Spacer(1, 48),
        text("DOCUMENTO DE PROCESSO E REQUISITOS PARA RPA", "kicker"),
        text(title_text, "cover"),
        text(f"Público: {job.get('audience') or 'Equipe de RPA'}  |  Nível: {job.get('detail_level') or 'operacional'}", "subtitle"),
        text(document.get("executive_summary"), "lead"),
        text("Este documento consolida o processo demonstrado. A transcrição e os prints são evidências de apoio e permanecem separados no pacote.", "note"),
        *([
            text("Guia recebido para esta análise", "h2"),
            text(job["process_context"]),
            text("Orientação da equipe para documentação e cortes; não equivale a um objetivo de negócio confirmado.", "note"),
        ] if job.get("process_context") else []),
        PageBreak(),
        text("1. Objetivo", "h1"),
        text(document.get("objective")),
        text("2. Escopo", "h1"),
        text("Incluído", "h2"),
        *bullets(document.get("scope", {}).get("in_scope", [])),
        text("Fora do escopo", "h2"),
        *bullets(document.get("scope", {}).get("out_of_scope", [])),
        text("3. Atores e responsabilidades", "h1"),
        *data_table(
            ["Ator", "Responsabilidade"],
            [[item.get("name"), item.get("responsibility")] for item in document.get("actors", [])] or [["Não identificado", "Validar com a área responsável"]],
            [1.7, 4.8],
        ),
        text("4. Sistemas envolvidos", "h1"),
        *data_table(
            ["Sistema", "Finalidade no processo"],
            [[item.get("name"), item.get("purpose")] for item in document.get("systems", [])] or [["Não identificado", "Validar durante o refinamento"]],
            [1.7, 4.8],
        ),
        text("5. Pré-requisitos", "h1"),
        *bullets(document.get("prerequisites", [])),
        text("6. Entradas e saídas", "h1"),
        text("Entradas", "h2"),
        *data_table(
            ["Entrada", "Origem", "Obrigatoriedade"],
            [[item.get("name"), item.get("source"), item.get("required")] for item in document.get("inputs", [])] or [["Não identificado", "A validar", "A validar"]],
            [2.5, 2.4, 1.6],
        ),
        text("Saídas", "h2"),
        *data_table(
            ["Saída", "Destino"],
            [[item.get("name"), item.get("destination")] for item in document.get("outputs", [])] or [["Não identificado", "A validar"]],
            [3.25, 3.25],
        ),
        text("7. Fluxo operacional", "h1"),
    ]

    for index, item in enumerate(document.get("process_flow", []), 1):
        refs = ", ".join(str(ref) for ref in item.get("evidence_refs", [])) or "sem referência visual direta"
        story.append(KeepTogether([
            text(f"{item.get('sequence') or index}. {item.get('title') or 'Etapa'}", "h2"),
            text(item.get("description")),
            text(item.get("actor"), label="Responsável:"),
            text(item.get("system"), label="Sistema:"),
            text(item.get("input"), label="Entrada:"),
            text(item.get("output"), label="Saída:"),
            *([text(item.get("decision_or_exception"), label="Decisão ou exceção:")] if item.get("decision_or_exception") else []),
            text(f"Evidências: {refs}", "evidence"),
        ]))

    story.append(text("8. Regras de negócio", "h1"))
    story += bullets([item.get("rule") for item in document.get("business_rules", []) if item.get("rule")], "Nenhuma regra foi explicitada na gravação.")
    story.append(text("9. Exceções e tratamentos", "h1"))
    for item in document.get("exceptions", []):
        story.append(KeepTogether([
            text(item.get("scenario"), label="Cenário:"),
            text(item.get("handling"), label="Tratamento:"),
            text(item.get("status"), label="Status:"),
        ]))
    if not document.get("exceptions"):
        story.append(text("Nenhuma exceção foi explicitada na gravação."))
    story.append(text("10. Riscos e controles", "h1"))
    for item in document.get("risks_and_controls", []):
        story.append(KeepTogether([
            text(item.get("risk"), label="Risco:"),
            text(item.get("impact"), label="Impacto:"),
            text(item.get("control"), label="Controle:"),
        ]))
    if not document.get("risks_and_controls"):
        story.append(text("Nenhum risco ou controle foi explicitado na gravação."))
    story.append(text("11. Oportunidades de automação", "h1"))
    for item in document.get("automation_opportunities", []):
        story.append(KeepTogether([
            text(item.get("opportunity") or "Oportunidade", "h2"),
            text(item.get("benefit"), label="Benefício:"),
            text(item.get("dependency"), label="Dependência:"),
        ]))
    if not document.get("automation_opportunities"):
        story.append(text("Nenhuma oportunidade foi confirmada com os insumos disponíveis."))
    story.append(text("12. Pontos a validar", "h1"))
    story += bullets(document.get("open_questions", []), "Nenhum ponto adicional registrado.")
    # Keep a short warning list with its heading, so the last caveat cannot
    # become an isolated final page. Long lists can still split naturally.
    story.append(KeepTogether([
        text("13. Limitações da análise", "h1"),
        *bullets(document.get("limitations", [])),
        text(EXPORT_TYPOGRAPHY_NOTE, "note"),
    ]))

    SimpleDocTemplate(
        str(path), pagesize=LETTER, title=title_text, author="Btime RPA Docs",
        subject="Documentação pós-discovery de RPA", leftMargin=inch, rightMargin=inch,
        topMargin=inch, bottomMargin=inch,
    ).build(story, onFirstPage=decorate, onLaterPages=decorate)
    (workspace / "documentacao-pdf-aviso.txt").unlink(missing_ok=True)
    return path


def write_structured_files(job: dict, document: dict, steps: list[dict], workspace: Path):
    def bullets(items: list, empty: str = "Não identificado na gravação.") -> str:
        return "\n".join(f"- {item}" for item in items) if items else f"- {empty}"

    guide = str(job.get("process_context") or "Não informado.")
    limitations = bullets(document.get("limitations", []), "Nenhuma limitação adicional registrada; validar com o responsável antes do uso operacional.")
    provenance = f'''## Guia recebido para esta análise

{guide}

Orientação da equipe para documentação e cortes; não equivale a um objetivo de negócio confirmado.

## Limitações da análise

{limitations}
'''

    flow_rows = "\n".join(
        f"| {item.get('sequence', index)} | {markdown_value(item.get('title'))} | {markdown_value(item.get('description'))} | {markdown_value(item.get('actor'))} | {markdown_value(item.get('system'))} |"
        for index, item in enumerate(document.get("process_flow", []), 1)
    ) or "| — | Não identificado | — | — | — |"
    procedure = f'''# Procedimento Operacional — {job.get("title", "Processo")}

## Objetivo

{document.get("objective") or "Não identificado na gravação."}

## Visão geral

{document.get("executive_summary") or "Não identificado na gravação."}

## Pré-requisitos

{bullets(document.get("prerequisites", []))}

## Fluxo operacional

| Etapa | Atividade | Como é realizada | Responsável | Sistema |
|---:|---|---|---|---|
{flow_rows}

## Exceções conhecidas

{bullets([f"**{item.get('scenario', 'Cenário')}** — {item.get('handling') or 'Tratamento não identificado'} ({item.get('status') or 'a validar'})" for item in document.get('exceptions', [])])}

## Pontos a validar

{bullets(document.get("open_questions", []), "Nenhum ponto adicional registrado.")}

{provenance}
'''
    (workspace / "procedimento-operacional.md").write_text(procedure, encoding="utf-8")

    actors = bullets([f"**{item.get('name', 'Não identificado')}** — {item.get('responsibility') or 'Responsabilidade não identificada'}" for item in document.get("actors", [])])
    systems = bullets([f"**{item.get('name', 'Não identificado')}** — {item.get('purpose') or 'Finalidade não identificada'}" for item in document.get("systems", [])])
    rules = bullets([item.get("rule") or "Regra não descrita" for item in document.get("business_rules", [])])
    risks = bullets([f"**Risco:** {item.get('risk') or 'Não identificado'} | **Impacto:** {item.get('impact') or 'Não identificado'} | **Controle:** {item.get('control') or 'Não identificado'}" for item in document.get("risks_and_controls", [])])
    opportunities = bullets([f"**{item.get('opportunity', 'Oportunidade')}** — Benefício: {item.get('benefit') or 'a validar'}; dependência: {item.get('dependency') or 'a validar'}" for item in document.get("automation_opportunities", [])])
    requirements = f'''# Requisitos para RPA — {job.get("title", "Processo")}

## Escopo incluído

{bullets(document.get("scope", {}).get("in_scope", []))}

## Fora do escopo

{bullets(document.get("scope", {}).get("out_of_scope", []))}

## Atores e responsabilidades

{actors}

## Sistemas envolvidos

{systems}

## Regras de negócio

{rules}

## Riscos e controles

{risks}

## Oportunidades de automação

{opportunities}

## Dúvidas e dependências para refinamento

{bullets(document.get("open_questions", []))}

{provenance}
'''
    (workspace / "requisitos-rpa.md").write_text(requirements, encoding="utf-8")
    write_docx_document(job, document, workspace)
    write_pdf_document(job, document, workspace)

    with (workspace / "matriz-evidencias.csv").open("w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.writer(csv_file, delimiter=";")
        writer.writerow(["frame_index", "timecode", "etapa_observada", "sistema", "controle_ou_campo", "evidencia", "incerteza", "arquivo", "fala_relacionada"])
        for step in steps:
            writer.writerow([
                step.get("frame_index", ""), step.get("timecode", ""), step.get("title", ""),
                step.get("system", ""), step.get("field_or_control", ""), step.get("evidence", ""),
                step.get("uncertainty", ""), step.get("image", ""), step.get("transcript", ""),
            ])


def locate_in_preview(source_previews: list[dict] | None, moment: float, lead_in: float = 8.0) -> dict | None:
    """Posiciona um instante da gravação dentro do preview já cortado, para revisitar a fala daquele trecho."""
    try:
        moment = float(moment)
    except (TypeError, ValueError):
        return None
    for preview in source_previews or []:
        if not preview.get("created") or not preview.get("file"):
            continue
        offset = float(preview.get("global_offset") or 0.0)
        local = moment - offset
        if local < 0:
            continue
        elapsed = 0.0
        for item in preview.get("ranges") or []:
            start = float(item["start"])
            end = float(item["end"])
            if start <= local <= end:
                position = elapsed + (local - start)
                return {
                    "file": preview["file"],
                    "play_from": round(max(elapsed, position - lead_in), 3),
                    "play_to": round(elapsed + (end - start), 3),
                    "cut_start": offset + start,
                    "cut_end": offset + end,
                    "source_name": preview.get("source_name") or "",
                }
            elapsed += end - start
    return None


def report_styles() -> str:
    """Self-contained offline report: official local fonts, no external requests."""
    fonts = []
    font_dir = Path(__file__).parent / "static" / "btime" / "fontes"
    for name, weight in (("Semilight", "300 400"), ("Medium", "500"), ("Semibold", "600 800")):
        path = font_dir / f"Branding-{name}.woff"
        if path.is_file():
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            fonts.append(f"@font-face{{font-family:Branding;src:url(data:font/woff;base64,{encoded}) format('woff');font-weight:{weight};font-display:swap}}")
    return "\n".join(fonts) + '''
:root{--night:#060315;--violet:#571ee6;--lavender:#d5c5ff;--paper:#f7f6fb;--muted:#62596f;--divider:#e1ddeb}
*{box-sizing:border-box}body{font:17px/1.55 Branding,Arial,sans-serif;margin:0;background:var(--paper);color:var(--night)}
header{padding:46px max(5vw,24px);background:var(--night);color:white}header p{max-width:850px}
h1{font-size:clamp(30px,4vw,44px);line-height:1.15;letter-spacing:-.025em;max-width:1000px;overflow-wrap:anywhere}
h2{margin-top:0;line-height:1.25;font-size:25px}h3{margin:5px 0;line-height:1.3}main{max-width:1120px;margin:30px auto;padding:0 20px}
section{margin:22px 0}.panel{background:white;padding:24px;border:1px solid var(--divider);border-radius:12px;min-width:0;overflow-x:auto}
.summary{font-size:21px;max-width:900px}.request-guide p{white-space:pre-wrap}.request-guide small{color:var(--muted)}
.deliverables{display:flex;flex-wrap:wrap;gap:8px;margin-top:20px}a{color:var(--violet);text-underline-offset:3px}a:hover{text-decoration:underline}
a:focus-visible,summary:focus-visible,video:focus-visible{outline:3px solid var(--violet);outline-offset:4px}header a:focus-visible,.preview a:focus-visible{outline-color:var(--lavender)}
.deliverables a,.evidence-ref{font-weight:600;text-decoration:none}.deliverables a{display:inline-flex;align-items:center;min-height:44px;padding:9px 12px;background:var(--lavender);color:var(--night);border-radius:8px}
.preview{background:var(--night);color:white;padding:24px;border-radius:12px}.preview a{color:var(--lavender)}.preview article+article{margin-top:28px;padding-top:20px;border-top:1px solid #817990}
video{display:block;width:100%;max-height:620px;margin-top:14px;background:#000;border-radius:8px}.columns{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:18px}
.entity-list{list-style:none;padding:0}.entity-list li{display:flex;flex-direction:column;padding:10px 0;border-bottom:1px solid var(--divider)}
.flow-step{display:grid;grid-template-columns:48px minmax(0,1fr);gap:15px;background:white;margin:12px 0;padding:20px;border:1px solid var(--divider);border-radius:12px;border-left:4px solid var(--violet)}
.flow-step>b{display:grid;place-items:center;width:38px;height:38px;border-radius:8px;background:var(--violet);color:white}.flow-step footer{display:flex;flex-wrap:wrap;gap:7px;margin-top:12px}
.evidence-ref{display:inline-flex;align-items:center;min-height:44px;padding:5px 10px;background:#f0eaff;border-radius:8px;font-size:14px}.not-confirmed,.empty{color:var(--muted);font-style:italic}
table{width:100%;border-collapse:collapse;background:white}th,td{padding:11px;text-align:left;border-bottom:1px solid var(--divider);vertical-align:top;overflow-wrap:anywhere}th{background:#f0eaff}
.rules li{margin:9px 0}.opportunities{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(240px,100%),1fr));gap:12px}.opportunity{background:#f0eaff;padding:20px;border-radius:12px}
.evidence-card{display:grid;grid-template-columns:42% minmax(0,1fr);gap:22px;background:white;margin:15px 0;padding:20px;border:1px solid var(--divider);border-radius:12px}
.evidence-card img{width:100%;border-radius:8px;border:1px solid var(--divider)}.evidence-media{min-width:0}.evidence-media video{margin-top:0;max-height:none;border:1px solid var(--divider)}
.evidence-media small{display:block;margin-top:8px;color:var(--muted);font-size:13px}span,dt{color:var(--muted);font-size:14px;font-weight:600}
dl{display:grid;grid-template-columns:150px minmax(0,1fr);gap:7px 12px}dd{margin:0}details{margin-top:12px}summary{padding:10px 0;min-height:44px;cursor:pointer}p,li,dd,h3{overflow-wrap:anywhere}
@media(max-width:760px){.columns,.evidence-card{grid-template-columns:1fr}.flow-step{grid-template-columns:38px minmax(0,1fr);gap:12px;padding:16px}dl{grid-template-columns:1fr}main{padding:0 16px}.panel{padding:20px}header{padding:32px 20px}}
@media print{body{background:white;font-size:11pt}header{background:white;color:var(--night);padding:0 0 20px}.deliverables,.preview,video{display:none}main{max-width:none;margin:0;padding:0}.panel,.flow-step,.evidence-card{box-shadow:none;break-inside:avoid}.columns{display:block}a{color:inherit}}
'''


def write_report(job: dict, steps: list[dict], workspace: Path, video: Path, document: dict | None = None, source_previews: list[dict] | None = None):
    document = normalize_documentation(document or fallback_documentation(job, steps), job, steps)
    referenced_frames = set()
    for section in (document.get("process_flow", []), document.get("business_rules", [])):
        for item in section:
            if isinstance(item, dict):
                referenced_frames.update(str(ref) for ref in item.get("evidence_refs", []) if ref is not None)
    relevant_steps = [step for step in steps if str(step.get("frame_index")) in referenced_frames]
    if not relevant_steps:
        irrelevant_visual_terms = ("participante", "câmera", "videoconferência", "avatar", "sem mudança", "continuidade da fala")
        relevant_steps = [step for step in steps if not any(term in f"{step.get('title', '')} {step.get('action', '')}".lower() for term in irrelevant_visual_terms)]
    sanitized_steps = []
    for step in relevant_steps:
        display_step = dict(step)
        visual_context = f"{step.get('title', '')} {step.get('action', '')} {step.get('system', '')} {step.get('field_or_control', '')}".lower()
        if any(term in visual_context for term in ("participante", "câmera", "videoconferência", "feed de câmera", "avatar")):
            display_step.update({
                "title": "Explicação registrada em áudio",
                "action": "O trecho foi mantido pelo conteúdo falado associado. A imagem mostra apenas o contexto da reunião, sem ação operacional visível.",
                "system": "Reunião gravada",
                "field_or_control": "Nenhum controle operacional visível",
                "evidence": "Frame usado somente como referência temporal para a explicação falada.",
                "uncertainty": "A gravação não possui diarização confiável; nenhuma fala é atribuída nominalmente.",
            })
        sanitized_steps.append(display_step)
    relevant_steps = sanitized_steps
    report = {
        "title": job["title"], "source": str(video), "audience": job["audience"],
        "context": job["process_context"], "documentation": document, "evidence_steps": relevant_steps,
    }
    serialized = json.dumps(report, ensure_ascii=False, indent=2)
    (workspace / "relatorio.json").write_text(serialized, encoding="utf-8")
    (workspace / "documentacao-processo.json").write_text(serialized, encoding="utf-8")
    write_structured_files(job, document, relevant_steps, workspace)
    safe = lambda value: escape(str(value or ""))

    def list_html(items: list, empty: str = "Não identificado na gravação.") -> str:
        return f"<ul>{''.join(f'<li>{safe(item)}</li>' for item in items)}</ul>" if items else f'<p class="empty">{safe(empty)}</p>'

    evidence_by_ref = {str(step.get("frame_index")): step for step in relevant_steps if step.get("frame_index") is not None}

    def evidence_links(refs: list) -> str:
        links = []
        for ref in refs or []:
            step = evidence_by_ref.get(str(ref))
            if step:
                links.append(f'<a class="evidence-ref" href="#evidence-{safe(ref)}">Evidência {safe(ref)} · {safe(step.get("timecode"))}</a>')
        return "".join(links) or '<span class="not-confirmed">Sem referência visual direta</span>'

    flow_cards = "".join(f'''<article class="flow-step"><b>{safe(item.get('sequence') or index)}</b><div><span>{safe(item.get('actor') or 'Responsável não identificado')} · {safe(item.get('system') or 'Sistema não identificado')}</span><h3>{safe(item.get('title') or 'Etapa')}</h3><p>{safe(item.get('description') or 'Descrição não identificada.')}</p><dl><dt>Entrada</dt><dd>{safe(item.get('input') or 'Não identificada')}</dd><dt>Saída</dt><dd>{safe(item.get('output') or 'Não identificada')}</dd><dt>Decisão ou exceção</dt><dd>{safe(item.get('decision_or_exception') or 'Não identificada')}</dd></dl><footer>{evidence_links(item.get('evidence_refs', []))}</footer></div></article>''' for index, item in enumerate(document.get("process_flow", []), 1))

    input_rows = "".join(f"<tr><td>{safe(item.get('name'))}</td><td>{safe(item.get('source') or 'Não identificada')}</td><td>{safe(item.get('required') or 'A validar')}</td></tr>" for item in document.get("inputs", [])) or '<tr><td colspan="3">Não identificado na gravação.</td></tr>'
    output_rows = "".join(f"<tr><td>{safe(item.get('name'))}</td><td>{safe(item.get('destination') or 'Não identificado')}</td></tr>" for item in document.get("outputs", [])) or '<tr><td colspan="2">Não identificado na gravação.</td></tr>'
    rules = "".join(f"<li><p>{safe(item.get('rule'))}</p>{evidence_links(item.get('evidence_refs', []))}</li>" for item in document.get("business_rules", [])) or '<li class="empty">Nenhuma regra foi explicitada.</li>'
    exceptions = "".join(f"<tr><td>{safe(item.get('scenario'))}</td><td>{safe(item.get('handling') or 'Não identificado')}</td><td>{safe(item.get('status') or 'A validar')}</td></tr>" for item in document.get("exceptions", [])) or '<tr><td colspan="3">Nenhuma exceção foi explicitada.</td></tr>'
    risks = "".join(f"<tr><td>{safe(item.get('risk'))}</td><td>{safe(item.get('impact') or 'Não identificado')}</td><td>{safe(item.get('control') or 'Não identificado')}</td></tr>" for item in document.get("risks_and_controls", [])) or '<tr><td colspan="3">Nenhum risco ou controle foi explicitado.</td></tr>'
    actors = "".join(f"<li><strong>{safe(item.get('name'))}</strong><span>{safe(item.get('responsibility') or 'Responsabilidade não identificada')}</span></li>" for item in document.get("actors", [])) or '<li class="empty">Não identificados.</li>'
    systems = "".join(f"<li><strong>{safe(item.get('name'))}</strong><span>{safe(item.get('purpose') or 'Finalidade não identificada')}</span></li>" for item in document.get("systems", [])) or '<li class="empty">Não identificados.</li>'
    opportunities = "".join(f'''<article class="opportunity"><h3>{safe(item.get('opportunity'))}</h3><p><b>Benefício:</b> {safe(item.get('benefit') or 'A validar')}</p><p><b>Dependência:</b> {safe(item.get('dependency') or 'A validar')}</p></article>''' for item in document.get("automation_opportunities", [])) or '<p class="empty">Nenhuma oportunidade foi confirmada com os insumos disponíveis.</p>'
    def evidence_media(step: dict) -> str:
        image = f'<img src="evidence/{safe(step.get("image"))}" alt="Evidência da etapa">'
        cut = locate_in_preview(source_previews, step.get("time"))
        if not cut:
            return f'<div class="evidence-media">{image}<small>Momento fora dos cortes do preview; a imagem permanece como referência.</small></div>'
        label = f'Corte {format_time(cut["cut_start"])}–{format_time(cut["cut_end"])}'
        if cut["source_name"]:
            label += f' · {cut["source_name"]}'
        return (
            f'<div class="evidence-media"><video controls preload="none" poster="evidence/{safe(step.get("image"))}"'
            f' src="{safe(cut["file"])}#t={cut["play_from"]},{cut["play_to"]}"></video>'
            f'<small>{safe(label)}</small></div>'
        )

    evidence_cards = "".join(f'''<article class="evidence-card" id="evidence-{safe(step.get('frame_index') or index)}">{evidence_media(step)}<div><span>{safe(step.get('timecode'))} · {safe(step.get('system') or 'Sistema não identificado')}</span><h3>{safe(step.get('title') or 'Etapa observada')}</h3><p>{safe(step.get('action'))}</p><dl><dt>Controle ou campo</dt><dd>{safe(step.get('field_or_control') or 'Não identificado')}</dd><dt>O que comprova</dt><dd>{safe(step.get('evidence'))}</dd></dl><details><summary>Ver fala relacionada</summary><p>{safe(step.get('transcript') or 'Sem fala detectada neste trecho.')}</p></details></div></article>''' for index, step in enumerate(relevant_steps, 1))
    preview_items = [item for item in (source_previews or []) if item.get("created") and item.get("file")]
    if preview_items:
        preview_cards = "".join(f'''<article><h3>{safe(item.get('source_name') or f"Gravação {index}")}</h3><p>{safe(format_time(item.get('duration', 0)))} de conteúdo relevante</p><video controls preload="metadata" src="{safe(item['file'])}"></video><p><a href="{safe(item['file'])}" target="_blank">Abrir preview</a> · <a href="{safe(item['file'])}" download>Baixar separadamente</a></p></article>''' for index, item in enumerate(preview_items, 1))
        preview = f'<section class="preview"><h2>Previews separados por arquivo</h2><p>Cada gravação mantém seus próprios cortes, guiados pelo conteúdo falado e pelo contexto informado.</p>{preview_cards}</section>'
    elif (workspace / "preview-processo.mp4").exists():
        preview = '<section class="preview"><h2>Preview do processo</h2><p>Trechos selecionados principalmente pelo conteúdo falado; os frames confirmam e ilustram cada explicação relevante.</p><video controls preload="metadata" src="preview-processo.mp4"></video><p><a href="preview-processo.mp4" target="_blank">Abrir preview em nova aba</a> · <a href="preview-processo.mp4" download>Baixar preview</a></p></section>'
    else:
        preview = ""
    guide_panel = (
        '<section class="panel request-guide"><h2>Guia recebido para esta análise</h2>'
        f'<p>{safe(job.get("process_context") or "Não informado.")}</p>'
        '<small>Orientação da equipe para documentação e cortes; não equivale a um objetivo de negócio confirmado.</small></section>'
    )
    html = f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{safe(job['title'])}</title><style>{report_styles()}</style></head><body><header><small>DOCUMENTAÇÃO ESTRUTURADA PÓS-DISCOVERY DE RPA</small><h1>{safe(job['title'])}</h1><p>Público: {safe(job.get('audience'))} · Nível: {safe(job.get('detail_level'))}</p><nav class="deliverables"><a href="documentacao-processo.docx" download>Documento Word editável</a><a href="documentacao-processo.pdf" download>Documento PDF</a><a href="procedimento-operacional.md" download>Procedimento operacional</a><a href="requisitos-rpa.md" download>Requisitos para RPA</a><a href="matriz-evidencias.csv" download>Matriz de evidências</a><a href="documentacao-processo.json" download>Dados estruturados</a></nav></header><main>{guide_panel}<section class="panel"><small>VISÃO GERAL</small><h2>Resumo executivo</h2><p class="summary">{safe(document.get('executive_summary'))}</p><h3>Objetivo</h3><p>{safe(document.get('objective'))}</p></section><section class="columns"><div class="panel"><h2>Escopo incluído</h2>{list_html(document.get('scope', {}).get('in_scope', []))}</div><div class="panel"><h2>Fora do escopo</h2>{list_html(document.get('scope', {}).get('out_of_scope', []))}</div></section><section class="columns"><div class="panel"><h2>Atores e responsabilidades</h2><ul class="entity-list">{actors}</ul></div><div class="panel"><h2>Sistemas envolvidos</h2><ul class="entity-list">{systems}</ul></div></section><section class="panel"><h2>Pré-requisitos</h2>{list_html(document.get('prerequisites', []))}</section><section class="columns"><div class="panel"><h2>Entradas</h2><table><thead><tr><th>Entrada</th><th>Origem</th><th>Obrigatoriedade</th></tr></thead><tbody>{input_rows}</tbody></table></div><div class="panel"><h2>Saídas</h2><table><thead><tr><th>Saída</th><th>Destino</th></tr></thead><tbody>{output_rows}</tbody></table></div></section><section><small>PROCESSO ATUAL</small><h2>Fluxo operacional consolidado</h2>{flow_cards or '<p class="empty">Nenhuma etapa foi identificada.</p>'}</section><section class="panel"><h2>Regras de negócio</h2><ol class="rules">{rules}</ol></section><section class="panel"><h2>Exceções e tratamentos</h2><table><thead><tr><th>Cenário</th><th>Tratamento</th><th>Status</th></tr></thead><tbody>{exceptions}</tbody></table></section><section class="panel"><h2>Riscos e controles</h2><table><thead><tr><th>Risco</th><th>Impacto</th><th>Controle</th></tr></thead><tbody>{risks}</tbody></table></section><section><h2>Oportunidades de automação</h2><div class="opportunities">{opportunities}</div></section><section class="columns"><div class="panel"><h2>Pontos a validar</h2>{list_html(document.get('open_questions', []), 'Nenhum ponto adicional registrado.')}</div><div class="panel"><h2>Limitações desta análise</h2>{list_html(document.get('limitations', []))}</div></section>{preview}<section><small>ANEXO DE COMPROVAÇÃO</small><h2>Matriz visual e falas relacionadas</h2><p>Cada evidência abre o corte em que o assunto é falado, começando pouco antes do print e seguindo até o fim do trecho. As evidências sustentam a documentação e não substituem o fluxo consolidado acima.</p>{evidence_cards or '<p>Nenhuma evidência visual foi identificada.</p>'}</section></main></body></html>'''
    (workspace / "relatorio.html").write_text(html, encoding="utf-8")


def process(job: dict, workspace: Path, progress):
    progress(5, "Organizando as gravações recebidas")
    video, source_videos = prepare_source(Path(job["source_path"]), workspace)
    progress(10, f"{len(source_videos)} gravação(ões) pronta(s) para análise")
    evidence = extract_evidence(video, workspace / "evidence")
    progress(35, f"{len(evidence)} evidências visuais capturadas")
    transcript = transcribe(video, workspace)
    progress(60, "Cruzando telas e transcrição")
    steps = analyze(evidence, transcript, workspace, job["process_context"], job["detail_level"])
    (workspace / "analise-visual.json").write_text(json.dumps(steps, ensure_ascii=False, indent=2), encoding="utf-8")
    progress(72, "Estruturando o processo, regras e requisitos de RPA")
    documentation = synthesize_documentation(job, steps, transcript, workspace)
    if not transcript:
        documentation["limitations"].append("Não foi possível obter falas da gravação; a análise contextual por áudio precisa ser validada antes do uso operacional.")
        documentation["preview_moments"] = []
    progress(80, "Montando previews separados com os trechos relevantes" if len(source_videos) > 1 else "Montando o preview com os trechos relevantes")
    preview, source_previews = create_job_previews(video, source_videos, steps, workspace, documentation.get("preview_moments", []))
    progress(92, "Montando a documentação e o pacote completo")
    write_report(job, steps, workspace, video, documentation, source_previews)
    create_package(workspace)
    warnings = [path.read_text(encoding="utf-8") for path in (workspace / "transcricao-aviso.txt", workspace / "documentacao-aviso.txt", workspace / "documentacao-pdf-aviso.txt") if path.is_file()]
    if not preview.get("created"):
        warnings.append("Nenhum preview contextual foi gerado. Consulte as limitações da documentação e revise o contexto ou a transcrição.")
    return {"video": str(video), "source_files": [str(item) for item in source_videos], "source_count": len(source_videos), "evidence_count": len(evidence), "transcript_segments": len(transcript), "step_count": len(steps), "process_step_count": len(documentation.get("process_flow", [])), "deliverables": ["relatorio.html", "documentacao-processo.docx", "documentacao-processo.pdf", "procedimento-operacional.md", "requisitos-rpa.md", "matriz-evidencias.csv", "documentacao-processo.json", "previews separados", "roteiro-cortes.json"], "preview": preview, "previews": source_previews, "warnings": warnings}
